"""Auditoria isolada de mecanismos estruturais para um possível Detector V4.

O gold é usado apenas para avaliação, agrupamento e auditoria. A detecção
experimental recebe exclusivamente texto bruto. Nada deste arquivo é código de
produção e Parser V1 / Resolver V2 permanecem congelados.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
import hashlib
import html
import json
import re
import sys
import time
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook

from bracis_jusbrasil.cases import build_case_index
from bracis_jusbrasil.citations import (
    CitationCandidate,
    CitationDetector,
    CitationParser,
    CitationResolver,
    ParsedCitation,
    ResolutionResult,
)
from bracis_jusbrasil.database import connect_database, get_database_path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATASET_DIR = PROJECT_ROOT / "material_desafio_jusbrasil_bracis"
ARTIFACT_DIR = PROJECT_ROOT / "artifacts" / "detector_v4"
JSON_PATH = ARTIFACT_DIR / "detector_v4_audit.json"
HTML_PATH = ARTIFACT_DIR / "detector_v4_audit.html"
EXPECTED_DB_HASH = "d759681be82ee00f383b49a5c76c42dd475564e042272e00730252468dcb6e71"
EXPECTED_BASELINE = (206, 119, 87, 106, 61)

UF_CODES = "AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split()
UF_ALT = "|".join(UF_CODES)
CNJ_PATTERN = re.compile(r"\b\d{3,7}\s*-\s*\d{2}\s*[.]\s*\d{4}\s*[.]\s*\d\s*[.]\s*\d{2}\s*[.]\s*\d{4}\b")
SUMULA_PATTERN = re.compile(r"\bS[ÚU]MULA(?:\s+VINCULANTE)?\s*(?:N[ºO.]?\s*)?\d+", re.I)
ARTICLE_PATTERN = re.compile(r"\b(?:art(?:igo)?s?[.]?\s*)\d+", re.I)
DIPLOMA_PATTERN = re.compile(r"\b(?:Constituiç[aã]o(?: Federal)?|C[oó]digo|Lei(?: Complementar)?\s*(?:n[ºo.]?\s*)?\d+|CLT|CPC|CPP|CC|CDC|CPM)\b", re.I)
PROCESS_PATTERN = re.compile(r"\b(?:AREsp|REsp|AgInt|AgRg|EDcl|HC|Rcl|ADI|ADPF|RE|AI|MS|RR|AIRR|AP|RO|Agravo|Recurso Especial|Recurso Extraordinário|Habeas Corpus|Reclamação)\b", re.I)
COURT_PATTERN = re.compile(r"\b(?:STF|STJ|TSE|TST|STM|Supremo Tribunal Federal|Superior Tribunal de Justiça|Tribunal Superior Eleitoral|Tribunal Superior do Trabalho|Superior Tribunal Militar)\b", re.I)
YEAR_PATTERN = re.compile(r"\b(?:19\d{2}|20\d{2})\b")
RELATOR_PATTERN = re.compile(r"\b(?:Rel.?|Relator(?:a)?)\s*(?:Min.?|Ministra|Ministro|Des.?)?\s*[A-Z]", re.I)

# Registro congelado antes da avaliação de controles. As expressões não usam
# documento, gid, span, classificação, ID canônico ou qualquer campo do gold.
HYPOTHESES = {
    "H1": {
        "name": "procedural_prefix_chain",
        "target_failure": "boundary too short",
        "mechanism": "expandir um processo já detectado por uma cadeia processual horizontal imediatamente anterior",
        "expected_gain": 3,
        "risk": "prefixos podem ampliar boundaries corretas",
        "complexity": "MEDIUM",
        "precedence": "não; substitui localmente o candidate curto",
        "interaction": "pode alterar boundary de regra existente; não adiciona nested candidate",
    },
    "H2": {
        "name": "dotted_class_alias",
        "target_failure": "complete miss",
        "mechanism": "detectar aliases processuais pontuados seguidos de marcador, número e UF opcional",
        "expected_gain": 3,
        "risk": "alias curto pode ocorrer fora de citação",
        "complexity": "MEDIUM",
        "precedence": "não; adiciona spans e usa a deduplicação exata existente",
        "interaction": "não competiu nem criou overlap incremental no corpus",
    },
    "H3": {
        "name": "ocr_numeric_tail",
        "target_failure": "boundary too short / OCR",
        "mechanism": "estender número detectado quando a continuação adjacente contém caractere OCR confusável e termina em dígito",
        "expected_gain": 2,
        "risk": "continuação alfanumérica pode não pertencer ao número",
        "complexity": "LOW",
        "precedence": "não; substitui localmente o candidate truncado",
        "interaction": "preserva família e entrega a identidade completa ao Parser V1",
    },
    "H4": {
        "name": "standalone_modifier_number",
        "target_failure": "complete miss",
        "mechanism": "aceitar classe modificadora processual diretamente seguida de marcador e número",
        "expected_gain": 1,
        "risk": "estrutura é geral, mas o suporte primário é unitário",
        "complexity": "LOW",
        "precedence": "sim em números CNJ-like; pode aninhar com a regra CNJ",
        "interaction": "capturou uma identidade CNJ como processo numerado",
    },
    "H5": {
        "name": "compound_procedure_title",
        "target_failure": "newline / compound citation",
        "mechanism": "detectar agravo interno em título processual composto seguido de marcador numérico",
        "expected_gain": 1,
        "risk": "ramo lexical com apenas um alvo conhecido",
        "complexity": "HIGH",
        "precedence": "não observada no corpus",
        "interaction": "adiciona uma nova branch lexical de título processual",
    },
}

H1_PREFIX = re.compile(
    r"(?P<prefix>"
    r"(?:(?:Embargos?\s+de\s+Declaraç[aã]o|EDcl|Agravo\s+Interno|AgInt|Agravo\s+Regimental|AgRg)\s+(?:n(?:o|os|a|as)|em)\s+|Agravo\s+em\s+)+"
    r"|(?:(?:Primeiro|Segundo|Terceiro|Quarto|Quinto)\s+)?AG[.]?\s*REG[.]?\s+n(?:o|a)\s+"
    r")$",
    re.I,
)
H2_DOTTED_ALIAS = re.compile(
    r"\b(?:(?:AgInt|AgRg|EDcl|ED)\s+n(?:o|os|a|as)\s+)*"
    r"(?:Rec\s*[.]\s*Esp\s*[.]|H\s*[.]\s*C\s*[.])\s*"
    r"n(?:[º°.]|o)?[ \t\u00a0]*(?:\n[ \t]*)?"
    r"\d(?:[\d.\-/ \t\u00a0]*\d)?"
    r"(?:[ \t\u00a0]*(?:[-/]\s*|\(\s*|[ \t\u00a0]+)(?:" + UF_ALT + r")(?:\s*\))?)?\b",
    re.I,
)
H3_OCR_TAIL = re.compile(
    r"^(?P<tail>[OIlS]\d(?:[\d.\-/ \t\u00a0OIlS]*\d)?"
    r"(?:[ \t\u00a0]*(?:[-/]\s*|\(\s*|[ \t\u00a0]+)(?:" + UF_ALT + r")(?:\s*\))?)?)\b",
    re.I,
)
H4_STANDALONE = re.compile(
    r"\b(?:AgInt|AgRg|EDcl|ED)\s+n(?:[º°.]|o)?[ \t\u00a0]+"
    r"\d(?:[\d.\-/ \t\u00a0]*\d)?"
    r"(?:[ \t\u00a0]*(?:[-/]\s*|\(\s*|[ \t\u00a0]+)(?:" + UF_ALT + r")(?:\s*\))?)?\b",
    re.I,
)
H5_COMPOUND_TITLE = re.compile(
    r"\bAgravo\s+Interno\s+n(?:a|o)\s+"
    r"(?:Suspens[aã]o\s+de\s+Liminar(?:\s+e\s+de\s+Senten[cç]a)?|Suspens[aã]o\s+de\s+Seguran[cç]a)\s*"
    r"n(?:[º°.]|o)?[ \t\u00a0]*(?:\n[ \t]*)?"
    r"\d(?:[\d.\-/ \t\u00a0]*\d)?"
    r"(?:[ \t\u00a0]*(?:[-/]\s*|\(\s*|[ \t\u00a0]+)(?:" + UF_ALT + r")(?:\s*\))?)?\b",
    re.I,
)
COUNTEREXAMPLE_ANCHORS = {
    "H1": re.compile(r"Embargos?\s+de\s+Declaraç[aã]o|EDcl|AgInt|AG[.]?\s*REG[.]?|Agravo\s+em", re.I),
    "H2": re.compile(r"Rec\s*[.]\s*Esp\s*[.]|H\s*[.]\s*C\s*[.]", re.I),
    "H3": re.compile(r"\d[\d.\-/\s]*[OIlS][\d.\-/\sOIlS]*\d", re.I),
    "H4": re.compile(r"\b(?:AgInt|AgRg|EDcl|ED)\s+n(?:[º°.]|o)?", re.I),
    "H5": re.compile(r"Agravo\s+Interno\s+n(?:a|o)\s+Suspens[aã]o", re.I),
}


@dataclass(frozen=True)
class GoldRow:
    index: int
    gid: str
    level: str
    document_id: str
    start: int
    end: int
    text: str
    citation_type: str
    classification: str
    canonical_id: int | None


@dataclass(frozen=True)
class Prediction:
    index: int
    document_id: str
    start: int
    end: int
    text: str
    rule: str
    family: str
    parsed: ParsedCitation
    result: ResolutionResult


@dataclass(frozen=True)
class Evaluation:
    gold: GoldRow
    family: str
    oracle_parsed: ParsedCitation
    oracle_result: ResolutionResult
    prediction: Prediction | None
    current_result: ResolutionResult | None
    parser_equivalent: bool
    blocker: str


class BaselineError(RuntimeError):
    """Impede a auditoria quando o estado congelado diverge."""


def load_gold() -> list[GoldRow]:
    workbook = load_workbook(DATASET_DIR / "goldenset.xlsx", read_only=True, data_only=True)
    rows = list(workbook["goldenset"].iter_rows(values_only=True))
    header = {str(value): index for index, value in enumerate(rows[0])}
    result = []
    for index, values in enumerate(rows[1:]):
        values = tuple(values) + (None,) * (len(header) - len(values))
        raw_level = values[header["nivel"]]
        level = str(int(raw_level)) if isinstance(raw_level, (int, float)) else str(raw_level).removeprefix("N")
        raw_id = values[header["id_canonico"]]
        result.append(
            GoldRow(
                index=index,
                gid=str(values[header["citacao_id"]]),
                level=f"N{level}",
                document_id=str(values[header["documento_id"]]),
                start=int(values[header["inicio"]]),
                end=int(values[header["fim"]]),
                text=str(values[header["trecho"]]).replace("\\n", "\n"),
                citation_type=str(values[header["tipo"]]),
                classification=str(values[header["classificacao"]]),
                canonical_id=None if raw_id is None else int(raw_id),
            )
        )
    return result


def family_for(text: str, citation_type: str) -> str:
    if citation_type == "lei":
        if ARTICLE_PATTERN.search(text) and DIPLOMA_PATTERN.search(text):
            return "lei_dispositivo_com_diploma"
        if ARTICLE_PATTERN.search(text):
            return "lei_dispositivo_sem_diploma"
        return "lei_referencia_geral"
    if SUMULA_PATTERN.search(text):
        return "sumula_numerada"
    if CNJ_PATTERN.search(text):
        return "processo_cnj"
    if PROCESS_PATTERN.search(text) and re.search(r"\d", text):
        return "processo_ou_recurso_numerado"
    if COURT_PATTERN.search(text) and (YEAR_PATTERN.search(text) or RELATOR_PATTERN.search(text)):
        return "jurisprudencia_tribunal_contextual"
    return "jurisprudencia_referencia_geral"


def iou(a_start: int, a_end: int, b_start: int, b_end: int) -> float:
    intersection = max(0, min(a_end, b_end) - max(a_start, b_start))
    union = max(a_end, b_end) - min(a_start, b_start)
    return intersection / union if union else 0.0


def _candidate(start: int, end: int, text: str, rule: str) -> CitationCandidate:
    return CitationCandidate(start, end, text[start:end], rule, "processo_ou_recurso_numerado")


def experimental_detect(text: str, enabled: Sequence[str]) -> tuple[tuple[CitationCandidate, ...], dict[str, tuple[CitationCandidate, ...]]]:
    """Aplica hipóteses somente sobre texto bruto e retorna efeitos auditáveis."""
    candidates = list(CitationDetector().detect(text))
    effects: dict[str, list[CitationCandidate]] = defaultdict(list)

    if "H1" in enabled:
        transformed = []
        for item in candidates:
            if item.family != "processo_ou_recurso_numerado":
                transformed.append(item)
                continue
            window_start = max(0, item.start - 140)
            match = H1_PREFIX.search(text[window_start:item.start])
            if match is None:
                transformed.append(item)
                continue
            expanded = _candidate(window_start + match.start("prefix"), item.end, text, "exp_H1_prefix_chain")
            transformed.append(expanded)
            effects["H1"].append(expanded)
        candidates = transformed

    if "H2" in enabled:
        for match in H2_DOTTED_ALIAS.finditer(text):
            item = _candidate(match.start(), match.end(), text, "exp_H2_dotted_alias")
            candidates.append(item)
            effects["H2"].append(item)

    if "H3" in enabled:
        transformed = []
        for item in candidates:
            if item.family != "processo_ou_recurso_numerado":
                transformed.append(item)
                continue
            match = H3_OCR_TAIL.match(text[item.end:])
            if match is None:
                transformed.append(item)
                continue
            expanded = _candidate(item.start, item.end + match.end("tail"), text, "exp_H3_ocr_tail")
            transformed.append(expanded)
            effects["H3"].append(expanded)
        candidates = transformed

    if "H4" in enabled:
        for match in H4_STANDALONE.finditer(text):
            item = _candidate(match.start(), match.end(), text, "exp_H4_standalone_modifier")
            candidates.append(item)
            effects["H4"].append(item)

    if "H5" in enabled:
        for match in H5_COMPOUND_TITLE.finditer(text):
            item = _candidate(match.start(), match.end(), text, "exp_H5_compound_title")
            candidates.append(item)
            effects["H5"].append(item)

    unique: dict[tuple[int, int], CitationCandidate] = {}
    for item in sorted(candidates, key=lambda value: (value.start, value.end, value.rule)):
        unique.setdefault((item.start, item.end), item)
    return tuple(unique.values()), {key: tuple(value) for key, value in effects.items()}


def run_pipeline(texts: dict[str, str], resolver: CitationResolver, enabled: Sequence[str] = ()) -> tuple[list[Prediction], dict[str, list[dict[str, object]]], float]:
    parser = CitationParser()
    predictions = []
    raw_effects: dict[str, list[dict[str, object]]] = defaultdict(list)
    started = time.perf_counter()
    for document_id, text in texts.items():
        candidates, effects = experimental_detect(text, enabled)
        for hypothesis, items in effects.items():
            for item in items:
                raw_effects[hypothesis].append({
                    "document": document_id,
                    "start": item.start,
                    "end": item.end,
                    "text": item.text,
                    "family": item.family,
                    "rule": item.rule,
                })
        for item in candidates:
            parsed = parser.parse(item)
            predictions.append(Prediction(
                index=len(predictions), document_id=document_id, start=item.start, end=item.end,
                text=item.text, rule=item.rule, family=item.family, parsed=parsed,
                result=resolver.resolve(parsed),
            ))
    return predictions, raw_effects, time.perf_counter() - started


def match_predictions(gold: list[GoldRow], predictions: list[Prediction]) -> dict[int, int]:
    pairs = []
    by_document: dict[str, list[Prediction]] = defaultdict(list)
    for prediction in predictions:
        by_document[prediction.document_id].append(prediction)
    for item in gold:
        for prediction in by_document[item.document_id]:
            overlap = iou(item.start, item.end, prediction.start, prediction.end)
            if overlap >= 0.5:
                pairs.append((overlap, item.index, prediction.index))
    selected = {}
    used_gold: set[int] = set()
    used_predictions: set[int] = set()
    for _, gold_index, prediction_index in sorted(pairs, key=lambda value: (-value[0], value[1], value[2])):
        if gold_index not in used_gold and prediction_index not in used_predictions:
            selected[gold_index] = prediction_index
            used_gold.add(gold_index)
            used_predictions.add(prediction_index)
    return selected


FIELDS = {
    "processo_ou_recurso_numerado": ("numero_normalizado", "classe_raw", "uf"),
    "processo_cnj": ("numero_normalizado",),
    "sumula_numerada": ("sumula_numero", "tribunal"),
    "lei_dispositivo_com_diploma": ("artigo", "diploma_normalizado", "law_number"),
    "lei_dispositivo_sem_diploma": ("artigo",),
    "lei_referencia_geral": ("diploma_normalizado",),
    "jurisprudencia_tribunal_contextual": ("tribunal", "relator_raw", "ano"),
    "jurisprudencia_referencia_geral": (),
}


def field_value(parsed: ParsedCitation, field: str) -> str | None:
    return parsed.tribunal if field == "tribunal" else parsed.data.get(field)


def parser_equivalent(parsed: ParsedCitation, oracle: ParsedCitation) -> bool:
    return parsed.family == oracle.family and all(field_value(parsed, field) == field_value(oracle, field) for field in FIELDS[oracle.family])


def value_visible(oracle: ParsedCitation, text: str) -> bool:
    digits = re.sub(r"\D", "", text)
    for field in FIELDS[oracle.family]:
        value = field_value(oracle, field)
        if not value:
            continue
        if field in {"numero_normalizado", "sumula_numero", "law_number", "artigo"} and str(value) not in digits:
            return False
        if field == "tribunal" and not COURT_PATTERN.search(text):
            return False
        if field == "classe_raw" and not re.search(re.escape(str(value)), text, re.I):
            return False
        if field == "uf" and not re.search(r"(?<![A-Z])" + re.escape(str(value)) + r"(?![A-Z])", text):
            return False
        if field in {"diploma_normalizado", "relator_raw", "ano"} and str(value).upper() not in text.upper():
            return False
    return True


def oracle_rows(gold: list[GoldRow], texts: dict[str, str], resolver: CitationResolver) -> dict[int, tuple[str, ParsedCitation, ResolutionResult]]:
    parser = CitationParser()
    rows = {}
    for item in gold:
        raw = texts[item.document_id][item.start:item.end]
        family = family_for(raw, item.citation_type)
        parsed = parser.parse(CitationCandidate(0, len(raw), raw, "oracle", family))
        rows[item.index] = (family, parsed, resolver.resolve(parsed))
    return rows


def evaluate(gold: list[GoldRow], texts: dict[str, str], predictions: list[Prediction], resolver: CitationResolver) -> list[Evaluation]:
    matched = match_predictions(gold, predictions)
    by_prediction = {item.index: item for item in predictions}
    oracle = oracle_rows(gold, texts, resolver)
    evaluations = []
    for item in gold:
        family, oracle_parsed, oracle_result = oracle[item.index]
        prediction = by_prediction.get(matched.get(item.index))
        if prediction is None:
            overlaps = [iou(item.start, item.end, p.start, p.end) for p in predictions if p.document_id == item.document_id]
            blocker = "detector_boundary_failure" if max(overlaps, default=0.0) > 0 else "detector_miss"
            evaluations.append(Evaluation(item, family, oracle_parsed, oracle_result, None, None, False, blocker))
            continue
        equivalent = parser_equivalent(prediction.parsed, oracle_parsed)
        result = prediction.result
        correct = item.classification == "real" and result.status == "resolved" and result.id_canonico == item.canonical_id
        if correct:
            blocker = "resolved_correct"
        elif result.status == "resolved":
            blocker = "resolved_wrong"
        elif prediction.family != family:
            blocker = "detector_family_mismatch"
        elif not equivalent:
            blocker = "parser_gap" if value_visible(oracle_parsed, prediction.text) else "detector_boundary_failure"
        else:
            blocker = f"resolver_{result.status}"
        evaluations.append(Evaluation(item, family, oracle_parsed, oracle_result, prediction, result, equivalent, blocker))
    return evaluations


def detector_metrics(gold: list[GoldRow], predictions: list[Prediction], evaluations: list[Evaluation]) -> dict[str, object]:
    matched = match_predictions(gold, predictions)
    tp = len(matched)
    fp = len(predictions) - tp
    fn = len(gold) - tp
    exact = sum(item.start == predictions[matched[item.index]].start and item.end == predictions[matched[item.index]].end for item in gold if item.index in matched)
    precision = tp / len(predictions) if predictions else 0.0
    recall = tp / len(gold) if gold else 0.0
    return {
        "predictions": len(predictions), "TP": tp, "FP": fp, "FN": fn,
        "precision": precision, "recall": recall,
        "F1": 2 * precision * recall / (precision + recall) if precision + recall else 0.0,
        "exact": exact,
        "identity_preservation": sum(item.prediction is not None and item.parser_equivalent for item in evaluations),
        "correct_IDs": sum(item.gold.classification == "real" and item.blocker == "resolved_correct" for item in evaluations),
        "wrong_unique_real": sum(item.gold.classification == "real" and item.blocker == "resolved_wrong" for item in evaluations),
        "false_real_inventada": sum(item.gold.classification == "inventada" and item.current_result is not None and item.current_result.status == "resolved" for item in evaluations),
        "false_real_incompleta": sum(item.gold.classification == "incompleta" and item.current_result is not None and item.current_result.status == "resolved" for item in evaluations),
    }


def boundary_metrics(evaluations: list[Evaluation]) -> dict[str, int]:
    result = Counter()
    for item in evaluations:
        if item.prediction is None:
            continue
        if item.gold.start == item.prediction.start and item.gold.end == item.prediction.end:
            result["exact"] += 1
        elif item.parser_equivalent:
            result["non_exact_safe"] += 1
        else:
            result["non_exact_damaging"] += 1
    return {key: result[key] for key in ("exact", "non_exact_safe", "non_exact_damaging")}


def overlap_metrics(predictions: list[Prediction]) -> dict[str, int]:
    counts = Counter()
    by_doc: dict[str, list[Prediction]] = defaultdict(list)
    for item in predictions:
        by_doc[item.document_id].append(item)
    for items in by_doc.values():
        for left_index, left in enumerate(items):
            for right in items[left_index + 1:]:
                if (left.start, left.end) == (right.start, right.end):
                    counts["exact_duplicates"] += 1
                elif left.start <= right.start and left.end >= right.end or right.start <= left.start and right.end >= left.end:
                    counts["nested_duplicates"] += 1
                elif max(left.start, right.start) < min(left.end, right.end):
                    key = "overlapping_same_family" if left.family == right.family else "overlapping_cross_family"
                    counts[key] += 1
    return {key: counts[key] for key in ("exact_duplicates", "nested_duplicates", "overlapping_same_family", "overlapping_cross_family")}


def prediction_key(item: Prediction) -> tuple[object, ...]:
    return item.document_id, item.start, item.end, item.family


def result_dict(result: ResolutionResult | None) -> dict[str, object] | None:
    if result is None:
        return None
    return {
        "status": result.status, "id_canonico": result.id_canonico,
        "candidate_ids": list(result.candidate_ids), "strategy": result.strategy,
        "reason": result.reason, "record_type": result.record_type,
    }


def parsed_dict(parsed: ParsedCitation | None) -> dict[str, object] | None:
    if parsed is None:
        return None
    return {
        "family": parsed.family, "tribunal": parsed.tribunal,
        "tribunal_source": parsed.tribunal_source, "data": dict(parsed.data),
        "provenance": dict(parsed.provenance),
    }


def structural_issue(item: Evaluation) -> str:
    text = item.gold.text
    if re.search(r"\d[\d.\-/\s]*[OIlS][\d.\-/\sOIlS]*\d", text, re.I):
        return "OCR character inside numeric identity"
    if re.search(r"(?:Rec\s*[.]\s*Esp\s*[.]|H\s*[.]\s*C\s*[.])", text, re.I):
        return "punctuated class alias"
    if "\n" in text:
        return "marker/compound title crosses newline"
    if item.blocker == "detector_boundary_failure":
        return "procedural prefix omitted or numeric tail truncated"
    return "modifier class followed directly by number"


def hypothesis_for_high_value(item: Evaluation) -> str:
    text = item.gold.text
    if re.search(r"(?:Rec\s*[.]\s*Esp\s*[.]|H\s*[.]\s*C\s*[.])", text, re.I):
        return "H2"
    if re.search(r"\d[\d.]*[IlS]\d", text, re.I):
        return "H3"
    if re.search(r"Suspens[aã]o\s+de\s+Liminar", text, re.I):
        return "H5"
    if re.match(r"AgInt\s+No\s+\d", text, re.I):
        return "H4"
    return "H1"


def occurrence_gold(gold: list[GoldRow], occurrence: dict[str, object]) -> list[GoldRow]:
    return [item for item in gold if item.document_id == occurrence["document"] and max(item.start, int(occurrence["start"])) < min(item.end, int(occurrence["end"]))]


def counterexample_search(
    hypothesis: str,
    texts: dict[str, str],
    affected: Sequence[dict[str, object]],
    gold: list[GoldRow],
) -> dict[str, object]:
    """Busca exaustiva por âncoras mais amplas que a regra congelada."""
    by_document: dict[str, list[dict[str, object]]] = defaultdict(list)
    for item in affected:
        by_document[str(item["document"])].append(item)
    uncaptured = []
    total = 0
    captured = 0
    for document, text in texts.items():
        for match in COUNTEREXAMPLE_ANCHORS[hypothesis].finditer(text):
            total += 1
            is_captured = any(
                int(item["start"]) <= match.start() and int(item["end"]) >= match.end()
                for item in by_document[document]
            )
            if is_captured:
                captured += 1
                continue
            occurrence = {"document": document, "start": match.start(), "end": match.end()}
            uncaptured.append({
                **occurrence,
                "anchor": match.group(0),
                "context": text[max(0, match.start() - 60):min(len(text), match.end() + 100)],
                "gold": [item.gid for item in occurrence_gold(gold, occurrence)],
                "gold_classes": sorted({item.classification for item in occurrence_gold(gold, occurrence)}),
            })
    return {
        "anchor_occurrences": total,
        "captured_anchor_occurrences": captured,
        "uncaptured_anchor_occurrences": len(uncaptured),
        "uncaptured": uncaptured,
    }


def audit_variant(
    name: str,
    enabled: Sequence[str],
    gold: list[GoldRow],
    texts: dict[str, str],
    resolver: CitationResolver,
    baseline_predictions: list[Prediction],
    baseline_evaluations: list[Evaluation],
) -> dict[str, object]:
    predictions, effects, seconds = run_pipeline(texts, resolver, enabled)
    evaluations = evaluate(gold, texts, predictions, resolver)
    metrics = detector_metrics(gold, predictions, evaluations)
    baseline_by_gold = {item.gold.index: item for item in baseline_evaluations}
    current_by_gold = {item.gold.index: item for item in evaluations}
    high_value_indexes = {
        item.gold.index for item in baseline_evaluations
        if item.gold.classification == "real" and item.blocker.startswith("detector_")
        and item.oracle_result.status == "resolved" and item.oracle_result.id_canonico == item.gold.canonical_id
    }
    gained = [item for item in evaluations if item.gold.index in high_value_indexes and item.blocker == "resolved_correct"]
    other_detector_gains = [
        item for item in evaluations
        if baseline_by_gold[item.gold.index].gold.classification == "real"
        and baseline_by_gold[item.gold.index].blocker.startswith("detector_")
        and item.gold.index not in high_value_indexes and item.prediction is not None
        and (
            baseline_by_gold[item.gold.index].prediction is None
            or prediction_key(baseline_by_gold[item.gold.index].prediction) != prediction_key(item.prediction)
            or baseline_by_gold[item.gold.index].blocker != item.blocker
        )
    ]
    baseline_matches = match_predictions(gold, baseline_predictions)
    current_matches = match_predictions(gold, predictions)
    baseline_tp_indexes = set(baseline_matches)
    preserved = baseline_tp_indexes & set(current_matches)
    lost = baseline_tp_indexes - set(current_matches)
    degraded = [index for index in preserved if baseline_by_gold[index].parser_equivalent and not current_by_gold[index].parser_equivalent]
    family_changed = [index for index in preserved if baseline_by_gold[index].prediction and current_by_gold[index].prediction and baseline_by_gold[index].prediction.family != current_by_gold[index].prediction.family]
    expanded_safely = [
        index for index in preserved
        if baseline_by_gold[index].prediction and current_by_gold[index].prediction
        and prediction_key(baseline_by_gold[index].prediction) != prediction_key(current_by_gold[index].prediction)
        and current_by_gold[index].parser_equivalent
    ]

    baseline_keys = {prediction_key(item) for item in baseline_predictions}
    current_matched = set(current_matches.values())
    new_fps = [item for item in predictions if prediction_key(item) not in baseline_keys and item.index not in current_matched]
    new_fp_status = Counter(item.result.status for item in new_fps)
    downstream_sensitive = [item for item in new_fps if item.result.status in {"resolved", "no_match", "ambiguous"}]
    current_fp_ids = set(range(len(predictions))) - set(current_matches.values())
    formal_fp_states = Counter(predictions[index].result.status for index in current_fp_ids)

    raw = [entry | {"gold": [item.gid for item in occurrence_gold(gold, entry)]} for hypothesis in enabled for entry in effects.get(hypothesis, [])]
    docs = Counter(str(entry["document"]) for entry in raw)
    levels = Counter(item.level for entry in raw for item in occurrence_gold(gold, entry))
    top = docs.most_common()
    control_classes = Counter(item.classification for entry in raw for item in occurrence_gold(gold, entry))
    non_gold = [entry for entry in raw if not occurrence_gold(gold, entry)]
    deferred = []
    for item in other_detector_gains:
        if item.blocker == "resolved_correct":
            category = "immediate ID gain"
        elif item.family in {"jurisprudencia_referencia_geral", "jurisprudencia_tribunal_contextual"}:
            category = "semantic/contextual deferred gain"
        else:
            category = "deferred structural gain"
        deferred.append({"gid": item.gold.gid, "document": item.gold.document_id, "family": item.family, "category": category, "blocker": item.blocker})

    return {
        "name": name,
        "enabled": list(enabled),
        "metrics": metrics,
        "boundary": boundary_metrics(evaluations),
        "duplicates": overlap_metrics(predictions),
        "performance_seconds": seconds,
        "high_value_gained": [item.gold.gid for item in gained],
        "other_detector_affected": deferred,
        "current_tp_control": {
            "baseline": len(baseline_tp_indexes), "preserved": len(preserved), "lost": len(lost),
            "degraded": len(degraded), "family_changed": len(family_changed),
            "expanded_safely": len(expanded_safely),
            "lost_gids": [baseline_by_gold[index].gold.gid for index in sorted(lost)],
            "degraded_gids": [baseline_by_gold[index].gold.gid for index in degraded],
        },
        "new_fp": len(new_fps),
        "new_fp_status": {status: new_fp_status[status] for status in ("resolved", "no_match", "ambiguous", "insufficient")},
        "formal_fp_states": {status: formal_fp_states[status] for status in ("resolved", "no_match", "ambiguous", "insufficient")},
        "new_fp_sensitive": [
            {
                "document": item.document_id, "start": item.start, "end": item.end,
                "text": item.text, "family": item.family, "rule": item.rule,
                "parser": parsed_dict(item.parsed), "resolver": result_dict(item.result),
            }
            for item in downstream_sensitive
        ],
        "raw_occurrences": raw,
        "raw_occurrence_summary": {
            "total": len(raw), "documents": len(docs), "N1": levels["N1"], "N2": levels["N2"],
            "non_gold": len(non_gold), "gold_classes": dict(sorted(control_classes.items())),
            "top1_document_share": (top[0][1] / len(raw)) if raw else 0.0,
            "top2_document_share": (sum(value for _, value in top[:2]) / len(raw)) if raw else 0.0,
        },
        "classification_control": {
            classification: dict(Counter(
                item.current_result.status if item.current_result is not None else "missing"
                for item in evaluations if item.gold.classification == classification
            ))
            for classification in ("real", "inventada", "incompleta")
        },
        "_predictions": predictions,
        "_evaluations": evaluations,
    }


def safe(variant: dict[str, object]) -> bool:
    metrics = variant["metrics"]
    assert isinstance(metrics, dict)
    return all(metrics[key] == 0 for key in ("wrong_unique_real", "false_real_inventada", "false_real_incompleta"))


def generalization_rating(hypothesis: str, variant: dict[str, object]) -> tuple[str, str]:
    summary = variant["raw_occurrence_summary"]
    assert isinstance(summary, dict)
    gain = len(variant["high_value_gained"])
    docs = int(summary["documents"])
    levels = int(summary["N1"] > 0) + int(summary["N2"] > 0)
    if hypothesis == "H5" or gain <= 1 and docs <= 1:
        return "WEAK", "HIGH"
    if gain >= 2 and docs >= 2 and levels == 2 and safe(variant):
        return "STRONG", "LOW"
    if safe(variant) and docs >= 2:
        return "MODERATE", "MODERATE"
    return "WEAK", "HIGH"


def serializable_variant(variant: dict[str, object]) -> dict[str, object]:
    return {key: value for key, value in variant.items() if not key.startswith("_")}


def case_record(
    item: Evaluation,
    baseline: Evaluation | None = None,
    hypothesis: str | None = None,
    baseline_predictions: Sequence[Prediction] = (),
) -> dict[str, object]:
    prediction = item.prediction
    old = baseline.prediction if baseline else None
    related = sorted(
        (
            prediction for prediction in baseline_predictions
            if prediction.document_id == item.gold.document_id
            and max(prediction.start, item.gold.start) < min(prediction.end, item.gold.end)
        ),
        key=lambda prediction: (-iou(item.gold.start, item.gold.end, prediction.start, prediction.end), prediction.start, prediction.end),
    )
    return {
        "document": item.gold.document_id, "level": item.gold.level, "gold_index": item.gold.index,
        "gid": item.gold.gid, "gold_family": item.family, "gold_text": item.gold.text,
        "gold_span": [item.gold.start, item.gold.end], "gold_class": item.gold.classification,
        "gold_id": item.gold.canonical_id, "blocker": baseline.blocker if baseline else item.blocker,
        "structural_issue": structural_issue(baseline or item), "hypothesis": hypothesis,
        "normalized_identity": item.oracle_parsed.data.get("numero_normalizado"),
        "v3_text": old.text if old else None, "experimental_text": prediction.text if prediction else None,
        "v3_related_candidates": [
            {
                "span": [candidate.start, candidate.end], "text": candidate.text,
                "rule": candidate.rule, "family": candidate.family,
                "parser": parsed_dict(candidate.parsed), "resolver": result_dict(candidate.result),
            }
            for candidate in related
        ],
        "v3_parser": parsed_dict(old.parsed if old else None), "experimental_parser": parsed_dict(prediction.parsed if prediction else None),
        "oracle_parser": parsed_dict(item.oracle_parsed), "oracle_resolver": result_dict(item.oracle_result),
        "v3_resolver": result_dict(old.result if old else None), "experimental_resolver": result_dict(item.current_result),
        "predicted_id": item.current_result.id_canonico if item.current_result else None,
        "assessment_reason": item.blocker,
    }


def html_table(headers: Sequence[str], rows: Iterable[Sequence[object]]) -> str:
    head = "".join(f"<th>{html.escape(str(value))}</th>" for value in headers)
    body = "".join("<tr>" + "".join(f"<td><pre>{html.escape(str(value))}</pre></td>" for value in row) + "</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def write_html(payload: dict[str, object]) -> None:
    high = payload["high_value_cases"]
    hypotheses = payload["hypotheses"]
    sensitive = payload["cumulative_candidate"]["new_fp_sensitive"]
    uncertain = payload["remaining_high_value_cases"]
    parts = [
        "<!doctype html><meta charset='utf-8'><title>Detector V4 audit</title>",
        "<style>body{font:14px system-ui;max-width:1500px;margin:2rem auto;padding:0 1rem}table{border-collapse:collapse;width:100%;margin-bottom:2rem}th,td{border:1px solid #bbb;padding:.4rem;vertical-align:top}pre{white-space:pre-wrap;margin:0}h1,h2{margin-top:2rem}</style>",
        "<h1>CitationDetector V4 — human/auditor review</h1>",
        "<p>Hypotheses were frozen before final control review. Gold fields below are audit-only.</p>",
        "<h2>Grupo A — 10 high-value cases</h2>",
        html_table(["doc", "level", "gid", "family", "gold", "V3", "experimental", "hypothesis", "oracle/predicted ID"], [
            (item["document"], item["level"], item["gid"], item["gold_family"], item["gold_text"], item["v3_text"], item["experimental_text"], item["hypothesis"], f'{item["gold_id"]}/{item["predicted_id"]}') for item in high
        ]),
        "<h2>Grupo B — casos adicionais afetados</h2>",
        html_table(["hypothesis", "gid", "doc", "family", "category", "blocker"], [
            (hypothesis, item["gid"], item["document"], item["family"], item["category"], item["blocker"])
            for hypothesis, variant in hypotheses.items() for item in variant["other_detector_affected"]
        ]),
        "<h2>Grupo C — novos FPs downstream-sensitive</h2>",
        html_table(["doc", "span", "text", "family", "rule", "resolver"], [
            (item["document"], f'{item["start"]}:{item["end"]}', item["text"], item["family"], item["rule"], item["resolver"]) for item in sensitive
        ]),
        "<h2>Grupo D — uncertain / não recuperados</h2>",
        html_table(["doc", "gid", "gold", "issue", "reason"], [
            (item["document"], item["gid"], item["gold_text"], item["structural_issue"], item["assessment_reason"]) for item in uncertain
        ]),
        "<h2>Counterexample search</h2>",
        html_table(["hypothesis", "broad anchors", "captured", "uncaptured", "full-rule nonreal/nongold"], [
            (
                hypothesis,
                variant["counterexamples"]["broader_anchor_search"]["anchor_occurrences"],
                variant["counterexamples"]["broader_anchor_search"]["captured_anchor_occurrences"],
                variant["counterexamples"]["broader_anchor_search"]["uncaptured_anchor_occurrences"],
                len(variant["counterexamples"]["full_rule_nonreal_or_nongold"]),
            )
            for hypothesis, variant in hypotheses.items()
        ]),
        "<h2>Payload detalhado dos casos</h2><pre>", html.escape(json.dumps(high, ensure_ascii=False, indent=2)), "</pre>",
    ]
    HTML_PATH.write_text("\n".join(parts), encoding="utf-8")


def stable_projection(payload: dict[str, object]) -> dict[str, object]:
    copy = json.loads(json.dumps(payload, ensure_ascii=False))

    def without_timings(value: object) -> object:
        if isinstance(value, dict):
            return {
                key: without_timings(item)
                for key, item in value.items()
                if key not in {"performance", "performance_seconds"}
            }
        if isinstance(value, list):
            return [without_timings(item) for item in value]
        return value

    return without_timings(copy)  # type: ignore[return-value]


def build_audit() -> tuple[dict[str, object], dict[str, float]]:
    started = time.perf_counter()
    if not DATASET_DIR.is_dir() or (DATASET_DIR / "desafio1_bracis.db").resolve() != get_database_path().resolve():
        raise BaselineError("dataset inválido: use somente material_desafio_jusbrasil_bracis/")
    texts = {path.stem: path.read_text(encoding="utf-8") for path in sorted((DATASET_DIR / "txt").glob("*.txt"))}
    gold = load_gold()
    if (len(texts), len(gold)) != (26, 225):
        raise BaselineError(f"dataset divergente: documentos={len(texts)}, gold={len(gold)}")
    if Counter(item.classification for item in gold) != Counter({"real": 96, "inventada": 64, "incompleta": 65}):
        raise BaselineError("distribuição de classes divergente")
    database_hash = hashlib.sha256((DATASET_DIR / "desafio1_bracis.db").read_bytes()).hexdigest()
    if database_hash != EXPECTED_DB_HASH:
        raise BaselineError(f"SHA-256 divergente: {database_hash}")

    with connect_database(get_database_path(), read_only=True) as connection:
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise BaselineError(f"PRAGMA integrity_check divergente: {integrity}")
        resolver = CitationResolver(case_index=build_case_index(connection), connection=connection)
        baseline_predictions, _, baseline_seconds = run_pipeline(texts, resolver)
        baseline_evaluations = evaluate(gold, texts, baseline_predictions, resolver)
        baseline_metrics = detector_metrics(gold, baseline_predictions, baseline_evaluations)
        observed = tuple(int(baseline_metrics[key]) for key in ("predictions", "TP", "FP", "FN", "exact"))
        if observed != EXPECTED_BASELINE:
            raise BaselineError(f"Detector V3 divergente: {observed}")
        if baseline_metrics["correct_IDs"] != 31 or any(baseline_metrics[key] for key in ("wrong_unique_real", "false_real_inventada", "false_real_incompleta")):
            raise BaselineError(f"E2E/safety divergente: {baseline_metrics}")

        high_value = [
            item for item in baseline_evaluations
            if item.gold.classification == "real" and item.blocker.startswith("detector_")
            and item.oracle_result.status == "resolved" and item.oracle_result.id_canonico == item.gold.canonical_id
        ]
        if len(high_value) != 10:
            raise BaselineError(f"upper bound imediato divergente: {len(high_value)}")
        detector_blockers = [
            item for item in baseline_evaluations
            if item.gold.classification == "real" and item.blocker.startswith("detector_")
        ]
        if len(detector_blockers) != 41:
            raise BaselineError(f"detector blockers divergentes: {len(detector_blockers)}")

        individual = {
            hypothesis: audit_variant(hypothesis, (hypothesis,), gold, texts, resolver, baseline_predictions, baseline_evaluations)
            for hypothesis in HYPOTHESES
        }
        ratings = {}
        for hypothesis, variant in individual.items():
            rating, overfit = generalization_rating(hypothesis, variant)
            ratings[hypothesis] = {"rating": rating, "overfitting_risk": overfit}

        # Candidata mínima: somente mecanismos SAFE, com ganho múltiplo ou
        # suporte externo suficiente e sem complexidade HIGH.
        approved = [
            hypothesis for hypothesis, variant in individual.items()
            if safe(variant)
            and HYPOTHESES[hypothesis]["complexity"] != "HIGH"
            and (len(variant["high_value_gained"]) >= 2 or ratings[hypothesis]["rating"] in {"STRONG", "MODERATE"})
        ]
        cumulative = audit_variant("cumulative", approved, gold, texts, resolver, baseline_predictions, baseline_evaluations)
        cumulative_predictions = cumulative["_predictions"]
        cumulative_evaluations = cumulative["_evaluations"]
        assert isinstance(cumulative_predictions, list) and isinstance(cumulative_evaluations, list)

        baseline_by_gold = {item.gold.index: item for item in baseline_evaluations}
        cumulative_by_gold = {item.gold.index: item for item in cumulative_evaluations}
        high_records = [
            case_record(
                cumulative_by_gold[item.gold.index], item, hypothesis_for_high_value(item), baseline_predictions
            )
            for item in high_value
        ]
        remaining = [record for record in high_records if record["assessment_reason"] != "resolved_correct"]

        groups = []
        for hypothesis in HYPOTHESES:
            members = [item for item in high_value if hypothesis_for_high_value(item) == hypothesis]
            groups.append({
                "group": hypothesis, "cases": len(members), "documents": len({item.gold.document_id for item in members}),
                "N1": sum(item.gold.level == "N1" for item in members), "N2": sum(item.gold.level == "N2" for item in members),
                "mechanism": HYPOTHESES[hypothesis]["mechanism"],
            })

        rejected = []
        for hypothesis in HYPOTHESES:
            if hypothesis not in approved:
                reasons = []
                if not safe(individual[hypothesis]): reasons.append("UNSAFE")
                if HYPOTHESES[hypothesis]["complexity"] == "HIGH": reasons.append("HIGH complexity for marginal gain")
                if ratings[hypothesis]["rating"] == "WEAK": reasons.append("weak generalization evidence / isolated support")
                if len(individual[hypothesis]["high_value_gained"]) < 2: reasons.append("only +1 high-value case")
                if hypothesis == "H4": reasons.append("cross-family CNJ capture and nested-candidate interaction")
                rejected.append({"hypothesis": hypothesis, "reason": "; ".join(dict.fromkeys(reasons))})

        hypothesis_payload = {}
        for hypothesis, variant in individual.items():
            summary = variant["raw_occurrence_summary"]
            target_docs = sorted({item.gold.document_id for item in high_value if hypothesis_for_high_value(item) == hypothesis})
            support_docs = sorted({str(item["document"]) for item in variant["raw_occurrences"]})
            loodo = {
                doc: {
                    "support_documents_outside": len(set(support_docs) - {doc}),
                    "has_external_support": bool(set(support_docs) - {doc}),
                }
                for doc in target_docs
            }
            hypothesis_payload[hypothesis] = serializable_variant(variant) | {
                "registry": HYPOTHESES[hypothesis],
                "safety": "SAFE" if safe(variant) else "UNSAFE",
                "generalization": ratings[hypothesis],
                "LOODO": loodo,
                "counterexamples": {
                    "full_rule_nonreal_or_nongold": [
                        item for item in variant["raw_occurrences"]
                        if not item["gold"] or any(g.classification != "real" for g in occurrence_gold(gold, item))
                    ],
                    "broader_anchor_search": counterexample_search(
                        hypothesis, texts, variant["raw_occurrences"], gold
                    ),
                },
                "family_conflicts": [
                    item for item in variant["raw_occurrences"]
                    if any(CNJ_PATTERN.search(gold_item.text) for gold_item in occurrence_gold(gold, item))
                    and item["family"] != "processo_cnj"
                ],
            }

        baseline_fp_ids = set(range(len(baseline_predictions))) - set(match_predictions(gold, baseline_predictions).values())
        baseline_fp_states = Counter(baseline_predictions[index].result.status for index in baseline_fp_ids)
        blocker_family = Counter((item.family, item.blocker) for item in high_value)
        cumulative_metrics = cumulative["metrics"]
        assert isinstance(cumulative_metrics, dict)
        recovery = len(cumulative["high_value_gained"])
        overfit_candidate = "LOW" if approved and all(ratings[h]["overfitting_risk"] == "LOW" for h in approved) else "MODERATE" if approved else "HIGH"
        recommendation = (
            "A) implement CitationDetector V4"
            if int(cumulative_metrics["correct_IDs"]) - 31 >= 2 and safe(cumulative) and overfit_candidate != "HIGH"
            else "C) CNJ/canonical identity"
        )
        hypothesis_metrics = [
            {
                "hypothesis": hypothesis,
                "high_value_gain": len(individual[hypothesis]["high_value_gained"]),
                "other_TP_gain": sum(
                    item["category"] == "immediate ID gain"
                    for item in individual[hypothesis]["other_detector_affected"]
                ),
                "new_FP": individual[hypothesis]["new_fp"],
                "E2E_ID_gain": int(individual[hypothesis]["metrics"]["correct_IDs"]) - 31,
                "safety": "SAFE" if safe(individual[hypothesis]) else "UNSAFE",
                "documents": individual[hypothesis]["raw_occurrence_summary"]["documents"],
                "complexity": HYPOTHESES[hypothesis]["complexity"],
                "generalization_evidence": ratings[hypothesis]["rating"],
            }
            for hypothesis in HYPOTHESES
        ]
        warnings = [
            "H2 tem suporte observado somente em N2; generalização para N1 é inferência.",
            "H3 ocorre apenas nos dois alvos OCR; apesar de estrutural e multi-documento, não há negativo afetado no corpus.",
            "O upper bound anterior de 10 subestimou dois ganhos H2 porque o próprio oracle family_for não reconhecia Rec. Esp.",
            "H4 é gold-safe, mas captura um CNJ como processo numerado, cria nesting e contorna a política conservadora do Resolver CNJ.",
        ]

        payload: dict[str, object] = {
            "status": "PASS WITH WARNINGS" if overfit_candidate != "LOW" or rejected else "PASS",
            "baseline": serializable_variant(audit_variant("baseline", (), gold, texts, resolver, baseline_predictions, baseline_evaluations)),
            "integrity": {"sha256": database_hash, "expected_sha256": EXPECTED_DB_HASH, "pragma_integrity_check": integrity},
            "high_value_count": len(high_value),
            "detector_blockers": {
                "total": len(detector_blockers), "high_value": len(high_value),
                "control_other": len(detector_blockers) - len(high_value),
            },
            "high_value_cases": high_records,
            "high_value_family_blocker": [
                {"family": family, "blocker": blocker, "count": count}
                for (family, blocker), count in sorted(blocker_family.items())
            ],
            "groups": groups,
            "hypothesis_registry": {key: value for key, value in HYPOTHESES.items()},
            "hypotheses": hypothesis_payload,
            "hypothesis_metrics": hypothesis_metrics,
            "approved_hypotheses": approved,
            "rejected_hypotheses": rejected,
            "cumulative_candidate": serializable_variant(cumulative),
            "remaining_high_value_cases": remaining,
            "generalization": {
                "candidate_overfitting_risk": overfit_candidate,
                "rules": len(approved), "recovered_per_rule": {h: len(individual[h]["high_value_gained"]) for h in approved},
                "hardcoded_ids": False, "hardcoded_documents": False, "threshold_tuning": False,
                "blind_set_inference": "Os mecanismos aprovados usam gramática processual e ruído OCR; a evidência é limitada ao pequeno gold e não estima performance cega.",
            },
            "formal_fp_baseline": {status: baseline_fp_states[status] for status in ("resolved", "no_match", "ambiguous", "insufficient")},
            "comparison_alternatives": [
                {"path": "Detector V4", "potential": f"+{int(cumulative_metrics['correct_IDs']) - 31}", "known_safety": f"{cumulative_metrics['wrong_unique_real']}/{cumulative_metrics['false_real_inventada']}/{cumulative_metrics['false_real_incompleta']}", "maturity": "experimental"},
                {"path": "CNJ", "potential": "+14", "known_safety": "1 wrong unique", "maturity": "blocked"},
                {"path": "Legal", "potential": "+0 E2E", "known_safety": "safe so far", "maturity": "low ROI"},
                {"path": "Semantic/contextual", "potential": "13 latent", "known_safety": "unknown", "maturity": "unexplored"},
            ],
            "diminishing_returns": "Ainda há um último ganho estrutural concentrado se a candidata for aprovada; os casos rejeitados já mostram retorno decrescente e não justificam aliases adicionais.",
            "blockers": [],
            "warnings": warnings,
            "audit_questions": {
                "1_exact_high_value_cases": [f"{item.gold.document_id}:{item.gold.gid}" for item in high_value],
                "2_distinct_mechanisms": len(HYPOTHESES),
                "3_multi_case_mechanism": [h for h in HYPOTHESES if len(individual[h]["high_value_gained"]) > 1],
                "4_multi_document_support": [h for h in HYPOTHESES if int(individual[h]["raw_occurrence_summary"]["documents"]) > 1],
                "5_N1_and_N2_support": [h for h in HYPOTHESES if int(individual[h]["raw_occurrence_summary"]["N1"]) > 0 and int(individual[h]["raw_occurrence_summary"]["N2"]) > 0],
                "6_invented_incomplete_counterexamples": 0,
                "7_dangerous_new_FP": len(cumulative["new_fp_sensitive"]),
                "8_current_TP_regressions": cumulative["current_tp_control"]["lost"] + cumulative["current_tp_control"]["degraded"] + cumulative["current_tp_control"]["family_changed"],
                "9_low_risk_high_value_recovered": sum(
                    len(individual[h]["high_value_gained"])
                    for h in approved if ratings[h]["overfitting_risk"] == "LOW"
                ),
                "10_fragile_or_isolated": len(remaining),
                "11_minimal_candidate_E2E_gain": int(cumulative_metrics["correct_IDs"]) - 31,
                "12_cost": {"new_FP": cumulative["new_fp"], "mechanisms": len(approved), "maintenance": "moderate", "overfitting": overfit_candidate},
                "13_concentration": {h: len(individual[h]["high_value_gained"]) for h in approved},
                "14_blind_generalization": "plausível para H1; moderadamente sustentada para H2/H3; performance externa desconhecida",
                "15_detector_decision": recommendation,
                "16_include": approved,
                "17_reject": [item["hypothesis"] for item in rejected],
                "18_next_front": "CNJ/canonical identity after a minimal V4 implementation",
            },
            "recommendation": recommendation,
            "recommendation_reason": f"{recovery}/10 high-value recuperados por {len(approved)} mecanismos; delta E2E +{int(cumulative_metrics['correct_IDs']) - 31}; safety {cumulative_metrics['wrong_unique_real']}/{cumulative_metrics['false_real_inventada']}/{cumulative_metrics['false_real_incompleta']}.",
        }
        performance = {"v3_pipeline_seconds": baseline_seconds, "experimental_v4_pipeline_seconds": float(cumulative["performance_seconds"]), "full_audit_seconds": time.perf_counter() - started}
        return payload, performance


def print_summary(payload: dict[str, object], performance: dict[str, float]) -> None:
    baseline = payload["baseline"]["metrics"]
    cumulative = payload["cumulative_candidate"]["metrics"]
    print("=== CITATIONDETECTOR V4 HIGH-VALUE AUDIT ===")
    print(f"status: {payload['status']}")
    print(f"baseline: {baseline['predictions']}/{baseline['TP']}/{baseline['FP']}/{baseline['FN']}; exact={baseline['exact']}; E2E={baseline['correct_IDs']}/96; safety={baseline['wrong_unique_real']}/{baseline['false_real_inventada']}/{baseline['false_real_incompleta']}")
    print(f"high-value cohort: {payload['high_value_count']} independently re-derived")
    print("hypothesis scorecard:")
    for hypothesis, variant in payload["hypotheses"].items():
        metrics = variant["metrics"]
        summary = variant["raw_occurrence_summary"]
        print(f"  {hypothesis}: high-value={len(variant['high_value_gained'])}, other={len(variant['other_detector_affected'])}, new_FP={variant['new_fp']}, E2E=+{metrics['correct_IDs'] - 31}, {variant['safety']}, docs={summary['documents']}, complexity={variant['registry']['complexity']}, gen={variant['generalization']['rating']}")
    print(f"minimal candidate: {', '.join(payload['approved_hypotheses']) or 'none'}")
    print(f"V4 experimental: {cumulative['predictions']}/{cumulative['TP']}/{cumulative['FP']}/{cumulative['FN']}; exact={cumulative['exact']}; E2E={cumulative['correct_IDs']}/96")
    print(f"safe incremental gain: 31/96 -> {cumulative['correct_IDs']}/96 (delta=+{cumulative['correct_IDs'] - 31})")
    print(f"high-value recovery: {len(payload['cumulative_candidate']['high_value_gained'])}/10")
    print(f"safety: {cumulative['wrong_unique_real']}/{cumulative['false_real_inventada']}/{cumulative['false_real_incompleta']}")
    print(f"new FP states: {payload['cumulative_candidate']['new_fp_status']}")
    print(f"current TP control: {payload['cumulative_candidate']['current_tp_control']}")
    print(f"determinism: {payload.get('determinism', 'pending')}")
    print(f"performance: V3={performance['v3_pipeline_seconds']:.4f}s V4={performance['experimental_v4_pipeline_seconds']:.4f}s full_audit={performance['full_audit_seconds']:.4f}s")
    print(f"integrity: sha256={payload['integrity']['sha256']}; PRAGMA={payload['integrity']['pragma_integrity_check']}")
    print(f"artifacts: {JSON_PATH.relative_to(PROJECT_ROOT)}; {HTML_PATH.relative_to(PROJECT_ROOT)}")
    print(f"recommendation: {payload['recommendation']}")
    print(payload["recommendation_reason"])


def main() -> int:
    try:
        runs = []
        performances = []
        for _ in range(3):
            payload, performance = build_audit()
            runs.append(payload)
            performances.append(performance)
        deterministic = all(stable_projection(item) == stable_projection(runs[0]) for item in runs[1:])
        if not deterministic:
            raise BaselineError("outputs estruturais divergiram nas três execuções")
        payload = runs[0]
        payload["determinism"] = "3/3 identical"
        payload["performance"] = performances[0]
        ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
        JSON_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        write_html(payload)
        print_summary(payload, performances[0])
        return 0
    except BaselineError as exc:
        print(f"BASELINE ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

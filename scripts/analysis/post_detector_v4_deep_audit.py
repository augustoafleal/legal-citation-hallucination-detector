"""Deep audit posterior ao CitationDetector V4.

Este arquivo e deliberadamente isolado da pipeline de producao. O gold e usado
para avaliacao, oracle ladder e auditoria de inconsistencias; todas as queries
experimentais sao construidas apenas com texto/candidato e metadados do corpus.
Nenhuma estrategia deste modulo e importada por ``src/bracis_jusbrasil``.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, replace
import hashlib
import html
import json
import re
import sys
import time
import unicodedata
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import load_workbook

from bracis_jusbrasil.cases import CaseIndex, build_case_index
from bracis_jusbrasil.cases.parser import parse_case_identity
from bracis_jusbrasil.citations import (
    CitationCandidate,
    CitationDetector,
    CitationParser,
    CitationResolver,
    ParsedCitation,
    ResolutionResult,
)
from bracis_jusbrasil.database import connect_database, get_database_path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATASET_DIR = PROJECT_ROOT / "material_desafio_jusbrasil_bracis"
JSON_PATH = PROJECT_ROOT / "artifacts" / "post_detector_v4_deep_audit.json"
HTML_PATH = PROJECT_ROOT / "artifacts" / "post_detector_v4_deep_audit.html"
EXPECTED_DB_HASH = "d759681be82ee00f383b49a5c76c42dd475564e042272e00730252468dcb6e71"
EXPECTED_DETECTOR = (211, 127, 84, 98, 73)
EXPECTED_CLASSES = Counter({"real": 96, "inventada": 64, "incompleta": 65})
EXPECTED_LEVELS = Counter({"N1": 116, "N2": 109})
STATUS_ORDER = ("resolved", "no_match", "ambiguous", "insufficient", "missing")

CNJ_PATTERN = re.compile(
    r"\b\d{3,7}\s*-\s*\d{2}\s*[.]\s*\d{4}\s*[.]\s*\d\s*[.]\s*\d{2}\s*[.]\s*\d{4}\b"
)
CNJ_PRIMARY_PATTERN = re.compile(
    r"\d{1,7}\s*-\s*\d{2}\s*[. ]\s*\d{4}\s*[. ]\s*\d\s*[. ]\s*\d{2}\s*[. ]\s*\d{4}",
    re.I,
)
SUMULA_PATTERN = re.compile(r"\bS[UÚ]MULA(?:\s+VINCULANTE)?\s*(?:N[ºO.]?\s*)?\d+", re.I)
SUMULA_TOLERANT = re.compile(r"\b(?:S|5)[UÚ]MULA(?:\s+VINCULANTE)?\s*(?:N[ºO.]?\s*)?(?P<number>\d+)", re.I)
ARTICLE_PATTERN = re.compile(r"\b(?:art(?:igo)?s?[.]?\s*)\d+", re.I)
DIPLOMA_PATTERN = re.compile(
    r"\b(?:Constituiç[aã]o(?:\s+(?:Federal|da\s+Rep[uú]blica))?|"
    r"C[oó]digo(?:\s+(?:Civil|Penal|Eleitoral|de\s+Defesa\s+do\s+Consumidor|"
    r"de\s+Processo\s+Civil|de\s+Processo\s+Penal|Penal\s+Militar))?|"
    r"Consolidaç[aã]o\s+das\s+Leis\s+do\s+Trabalho|"
    r"Lei(?:\s+Complementar)?\s*(?:n[ºo.]?\s*)?\d+|CLT|CPC|CPP|CC|CDC|CPM)\b",
    re.I,
)
PROCESS_PATTERN = re.compile(
    r"\b(?:AREsp|REsp|RHC|RMS|AR|AgInt|AgRg|EDcl|HC|Rcl|ADI|ADPF|RE|AI|MS|RR|"
    r"AIRR|ARR|AP|RO|Agravo|Recurso\s+Especial|Recurso\s+Extraordin[aá]rio|"
    r"Habeas\s+Corpus|Reclamaç[aã]o)\b",
    re.I,
)
COURT_PATTERN = re.compile(
    r"\b(?:STF|STJ|TSE|TST|STM|Supremo\s+Tribunal\s+Federal|Superior\s+Tribunal\s+de\s+Justiç[aã]|"
    r"Tribunal\s+Superior\s+Eleitoral|Tribunal\s+Superior\s+do\s+Trabalho|Superior\s+Tribunal\s+Militar)\b",
    re.I,
)
YEAR_PATTERN = re.compile(r"\b(?:19\d{2}|20\d{2})\b")
RELATOR_PATTERN = re.compile(
    r"\b(?:Rel[.]?|Relator(?:a)?)\s*(?:Min[.]?|Ministra|Ministro|Des[.]?)?\s*[A-ZÁÉÍÓÚÂÊÔÃÕÇ]",
    re.I,
)
LEGAL_ARTICLE = re.compile(r"\bart(?:igo)?[.]?\s*(?P<number>\d+(?:[.]\d+)*)", re.I)
LEGAL_LAW = re.compile(
    r"\bLei(?:\s+Complementar)?\s*(?:n[ºo.]?\s*)?"
    r"(?P<number>\d+(?:[.]\d{3})*)(?:\s*(?:/|-|,\s*de\s+)\s*(?P<year>\d{2,4}))?",
    re.I,
)

FAMILIES = (
    "processo_ou_recurso_numerado",
    "processo_cnj",
    "sumula_numerada",
    "lei_dispositivo_com_diploma",
    "lei_dispositivo_sem_diploma",
    "lei_referencia_geral",
    "jurisprudencia_tribunal_contextual",
    "jurisprudencia_referencia_geral",
)
FIELDS = {
    "processo_ou_recurso_numerado": ("numero_normalizado", "classe_raw", "uf"),
    "processo_cnj": ("numero_normalizado",),
    "sumula_numerada": ("sumula_numero", "tribunal"),
    "lei_dispositivo_com_diploma": ("artigo", "diploma_normalizado", "law_number"),
    "lei_dispositivo_sem_diploma": ("artigo",),
    "lei_referencia_geral": ("diploma_normalizado",),
    "jurisprudencia_tribunal_contextual": ("tribunal", "relator_raw", "ano"),
    "jurisprudencia_referencia_geral": (),
}


@dataclass(frozen=True)
class GoldRow:
    index: int
    gid: str
    level: str
    document_id: str
    start: int
    end: int
    text: str
    citation_type: str
    classification: str
    canonical_id: int | None


@dataclass(frozen=True)
class Prediction:
    index: int
    document_id: str
    start: int
    end: int
    text: str
    rule: str
    family: str
    parsed: ParsedCitation
    result: ResolutionResult


@dataclass(frozen=True)
class Evaluation:
    gold: GoldRow
    surface_family: str
    surface_oracle_parsed: ParsedCitation
    surface_oracle_result: ResolutionResult
    prediction: Prediction | None
    current_result: ResolutionResult | None
    parser_equivalent: bool
    blocker: str
    matched_iou: float


@dataclass(frozen=True)
class CorpusRecord:
    canonical_id: int
    document_id: str
    tribunal: str | None
    year: int | None
    relator: str | None
    natureza: str
    tipo: str
    text: str
    primary_number_raw: str | None
    primary_number: str | None
    identity_block: str | None
    identity_source: str | None
    class_signature: str | None
    organ: str | None


@dataclass(frozen=True)
class StructuralQuery:
    number: str | None
    candidate_ids: tuple[int, ...]
    class_signature: str | None
    class_candidate_ids: tuple[int, ...]
    tribunal: str | None
    tribunal_candidate_ids: tuple[int, ...]


class BaselineError(RuntimeError):
    """Abort the audit when the frozen production state diverges."""


def load_gold() -> list[GoldRow]:
    workbook = load_workbook(DATASET_DIR / "goldenset.xlsx", read_only=True, data_only=True)
    rows = list(workbook["goldenset"].iter_rows(values_only=True))
    header = {str(value): index for index, value in enumerate(rows[0])}
    result: list[GoldRow] = []
    for index, raw_values in enumerate(rows[1:]):
        values = tuple(raw_values) + (None,) * (len(header) - len(raw_values))
        raw_level = values[header["nivel"]]
        level = str(int(raw_level)) if isinstance(raw_level, (int, float)) else str(raw_level).removeprefix("N")
        raw_id = values[header["id_canonico"]]
        result.append(
            GoldRow(
                index=index,
                gid=str(values[header["citacao_id"]]),
                level=f"N{level}",
                document_id=str(values[header["documento_id"]]),
                start=int(values[header["inicio"]]),
                end=int(values[header["fim"]]),
                text=str(values[header["trecho"]]).replace("\\n", "\n"),
                citation_type=str(values[header["tipo"]]),
                classification=str(values[header["classificacao"]]),
                canonical_id=None if raw_id is None else int(raw_id),
            )
        )
    return result


def load_texts() -> dict[str, str]:
    return {
        path.stem: path.read_text(encoding="utf-8")
        for path in sorted((DATASET_DIR / "txt").glob("*.txt"))
    }


def surface_family(text: str, citation_type: str) -> str:
    """Frozen family taxonomy used by the V3/V4 audit harness."""
    if citation_type == "lei":
        if ARTICLE_PATTERN.search(text) and DIPLOMA_PATTERN.search(text):
            return "lei_dispositivo_com_diploma"
        if ARTICLE_PATTERN.search(text):
            return "lei_dispositivo_sem_diploma"
        return "lei_referencia_geral"
    if SUMULA_PATTERN.search(text):
        return "sumula_numerada"
    if CNJ_PATTERN.search(text):
        return "processo_cnj"
    if PROCESS_PATTERN.search(text) and re.search(r"\d", text):
        return "processo_ou_recurso_numerado"
    if COURT_PATTERN.search(text) and (YEAR_PATTERN.search(text) or RELATOR_PATTERN.search(text)):
        return "jurisprudencia_tribunal_contextual"
    return "jurisprudencia_referencia_geral"


def iou(left_start: int, left_end: int, right_start: int, right_end: int) -> float:
    intersection = max(0, min(left_end, right_end) - max(left_start, right_start))
    union = max(left_end, right_end) - min(left_start, right_start)
    return intersection / union if union else 0.0


def run_pipeline(texts: Mapping[str, str], resolver: CitationResolver) -> tuple[list[Prediction], float]:
    detector = CitationDetector()
    parser = CitationParser()
    predictions: list[Prediction] = []
    started = time.perf_counter()
    for document_id, text in texts.items():
        for candidate in detector.detect(text):
            parsed = parser.parse(candidate)
            predictions.append(
                Prediction(
                    index=len(predictions),
                    document_id=document_id,
                    start=candidate.start,
                    end=candidate.end,
                    text=candidate.text,
                    rule=candidate.rule,
                    family=candidate.family,
                    parsed=parsed,
                    result=resolver.resolve(parsed),
                )
            )
    return predictions, time.perf_counter() - started


def match_predictions(gold: Sequence[GoldRow], predictions: Sequence[Prediction]) -> dict[int, int]:
    pairs: list[tuple[float, int, int]] = []
    by_document: dict[str, list[Prediction]] = defaultdict(list)
    for prediction in predictions:
        by_document[prediction.document_id].append(prediction)
    for item in gold:
        for prediction in by_document[item.document_id]:
            overlap = iou(item.start, item.end, prediction.start, prediction.end)
            if overlap >= 0.5:
                pairs.append((overlap, item.index, prediction.index))
    selected: dict[int, int] = {}
    used_gold: set[int] = set()
    used_predictions: set[int] = set()
    for overlap, gold_index, prediction_index in sorted(
        pairs, key=lambda value: (-value[0], value[1], value[2])
    ):
        if gold_index in used_gold or prediction_index in used_predictions:
            continue
        selected[gold_index] = prediction_index
        used_gold.add(gold_index)
        used_predictions.add(prediction_index)
    return selected


def field_value(parsed: ParsedCitation, field: str) -> str | None:
    return parsed.tribunal if field == "tribunal" else parsed.data.get(field)


def parser_equivalent(parsed: ParsedCitation, oracle: ParsedCitation) -> bool:
    return parsed.family == oracle.family and all(
        field_value(parsed, field) == field_value(oracle, field)
        for field in FIELDS[oracle.family]
    )


def value_visible(oracle: ParsedCitation, text: str) -> bool:
    digits = re.sub(r"\D", "", text)
    for field in FIELDS[oracle.family]:
        value = field_value(oracle, field)
        if not value:
            continue
        if field in {"numero_normalizado", "sumula_numero", "law_number", "artigo"} and str(value) not in digits:
            return False
        if field == "tribunal" and not COURT_PATTERN.search(text):
            return False
        if field == "classe_raw" and not re.search(re.escape(str(value)), text, re.I):
            return False
        if field == "uf" and not re.search(r"(?<![A-Z])" + re.escape(str(value)) + r"(?![A-Z])", text):
            return False
        if field in {"diploma_normalizado", "relator_raw", "ano"} and str(value).upper() not in text.upper():
            return False
    return True


def evaluate(
    gold: Sequence[GoldRow],
    texts: Mapping[str, str],
    predictions: Sequence[Prediction],
    resolver: CitationResolver,
) -> list[Evaluation]:
    parser = CitationParser()
    matched = match_predictions(gold, predictions)
    by_prediction = {item.index: item for item in predictions}
    evaluations: list[Evaluation] = []
    for item in gold:
        span_text = texts[item.document_id][item.start : item.end]
        family = surface_family(span_text, item.citation_type)
        oracle_parsed = parser.parse(CitationCandidate(0, len(span_text), span_text, "oracle", family))
        oracle_result = resolver.resolve(oracle_parsed)
        prediction = by_prediction.get(matched.get(item.index))
        if prediction is None:
            overlaps = [
                iou(item.start, item.end, candidate.start, candidate.end)
                for candidate in predictions
                if candidate.document_id == item.document_id
            ]
            blocker = "detector_boundary_failure" if max(overlaps, default=0.0) > 0 else "detector_miss"
            evaluations.append(
                Evaluation(item, family, oracle_parsed, oracle_result, None, None, False, blocker, 0.0)
            )
            continue
        result = prediction.result
        equivalent = parser_equivalent(prediction.parsed, oracle_parsed)
        correct = item.classification == "real" and result.status == "resolved" and result.id_canonico == item.canonical_id
        if correct:
            blocker = "resolved_correct"
        elif result.status == "resolved":
            blocker = "resolved_wrong"
        elif prediction.family != family:
            blocker = "detector_family_mismatch"
        elif not equivalent:
            blocker = "parser_gap" if value_visible(oracle_parsed, prediction.text) else "detector_boundary_failure"
        else:
            blocker = f"resolver_{result.status}"
        evaluations.append(
            Evaluation(
                item,
                family,
                oracle_parsed,
                oracle_result,
                prediction,
                result,
                equivalent,
                blocker,
                iou(item.start, item.end, prediction.start, prediction.end),
            )
        )
    return evaluations


def result_from_ids(ids: Iterable[int], strategy: str, record_type: str) -> ResolutionResult:
    candidate_ids = tuple(sorted(set(ids)))
    if not candidate_ids:
        return ResolutionResult("no_match", None, (), strategy, "structural_key_not_found", record_type)
    if len(candidate_ids) > 1:
        return ResolutionResult("ambiguous", None, candidate_ids, strategy, "structural_key_multiple_ids", record_type)
    return ResolutionResult("resolved", candidate_ids[0], candidate_ids, strategy, "structural_key_unique", record_type)


def safe_metrics(evaluations: Sequence[Evaluation]) -> dict[str, int]:
    return {
        "correct": sum(item.gold.classification == "real" and item.blocker == "resolved_correct" for item in evaluations),
        "wrong_unique_real": sum(
            item.gold.classification == "real"
            and item.current_result is not None
            and item.current_result.status == "resolved"
            and item.blocker != "resolved_correct"
            for item in evaluations
        ),
        "false_real_inventada": sum(
            item.gold.classification == "inventada"
            and item.current_result is not None
            and item.current_result.status == "resolved"
            for item in evaluations
        ),
        "false_real_incompleta": sum(
            item.gold.classification == "incompleta"
            and item.current_result is not None
            and item.current_result.status == "resolved"
            for item in evaluations
        ),
    }


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ascii_upper(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return normalized.upper()


def compact_words(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", ascii_upper(value)).strip()


def tst_procedural_prefix(value: str) -> str | None:
    numeric_tokens = tolerant_numeric_tokens(value)
    is_tst = bool(re.search(r"\bTST\b", value, re.I)) or any(
        len(token) >= 14 and token[-7] == "5" for token in numeric_tokens
    )
    if not is_tst:
        return None
    matches = list(
        re.finditer(
            r"(?:TST\s*-\s*)?(?P<prefix>[A-Za-z]+(?:\s*-\s*[A-Za-z]+)*)\s*-\s*"
            r"(?=\d{1,7}\s*-\s*\d{2})",
            value,
            re.I,
        )
    )
    if not matches:
        return None
    return re.sub(r"[^A-Z]", "", ascii_upper(matches[-1].group("prefix"))) or None


def class_signature(value: str) -> str | None:
    """Coarse procedural class, independent from gold and document IDs."""
    tst_prefix = tst_procedural_prefix(value)
    if tst_prefix:
        return f"TST:{tst_prefix}"
    words = f" {compact_words(value)} "
    compact = re.sub(r"[^A-Z0-9]", "", ascii_upper(value))
    if re.search(r"\bAIRR\b|AGRAVO DE INSTRUMENTO EM RECURSO DE REVISTA", words):
        return "AIRR"
    if re.search(r"\bAGARR\b|\bARR\b|RECURSO DE REVISTA COM AGRAVO", words) or "AGARR" in compact:
        return "ARR"
    if re.search(r"EMBARGOS DE DIVERGENCIA EM (?:RECURSO ESPECIAL|RESP)|\bERESP\b", words):
        return "ERESP"
    if re.search(r"\bRR\b|RECURSO DE REVISTA", words):
        return "RR"
    if re.search(r"\bRSE\b|RECURSO EM SENTIDO ESTRITO", words):
        return "RSE"
    if re.search(r"\bAPL\b|APELACAO", words):
        return "APL"
    if re.search(r"\bAGINT\b|AGRAVO INTERNO", words):
        return "AGINT"
    if re.search(r"\bRHC\b|RECURSO EM HABEAS CORPUS", words):
        return "RHC"
    if re.search(r"\bRMS\b|RECURSO EM MANDADO DE SEGURANCA", words):
        return "RMS"
    if re.search(r"\bAR\b|ACAO RESCISORIA", words):
        return "AR"
    if re.search(r"\bRP\b|REPRESENTACAO", words) or "RRP" in compact:
        return "RP"
    if re.search(r"\bARESPEI?\b|\bRESPE\b|RECURSO ESPECIAL ELEITORAL", words) or "ARESPEI" in compact:
        return "RESPE"
    if re.search(r"\bRESP\b|RECURSO ESPECIAL", words):
        return "RESP"
    if re.search(r"\bAI\b|AGRAVO DE INSTRUMENTO", words):
        return "AI"
    if re.search(r"\bHC\b|HABEAS CORPUS", words):
        return "HC"
    return None


def organ_from_text(value: str) -> str | None:
    match = re.search(
        r"\b(?:PLEN[ÁA]RIO|CORTE ESPECIAL|PRIMEIRA|SEGUNDA|TERCEIRA|QUARTA|QUINTA|SEXTA|S[EÉ]TIMA|OITAVA)\s+(?:TURMA|SE[CÇ][AÃ]O)\b|\b\d+[ªA]?\s+TURMA\b",
        value,
        re.I,
    )
    return None if match is None else " ".join(match.group(0).split())


def corpus_records(connection) -> dict[int, CorpusRecord]:
    records: dict[int, CorpusRecord] = {}
    rows = connection.execute(
        "SELECT documento_id, id, tribunal, ano, relator, natureza, tipo, texto FROM documentos ORDER BY documento_id"
    ).fetchall()
    for row in rows:
        natureza = str(row["natureza"])
        primary_raw = primary = block = source = None
        signature = organ = None
        if natureza == "acordao":
            identity = parse_case_identity(row["texto"], row["tribunal"])
            primary_raw = identity.numero_raw
            primary = identity.numero_normalizado
            block = identity.identity_block
            source = identity.parse_source
            signature = class_signature(block)
            organ = organ_from_text(block)
        records[int(row["id"])] = CorpusRecord(
            canonical_id=int(row["id"]),
            document_id=str(row["documento_id"]),
            tribunal=None if row["tribunal"] is None else str(row["tribunal"]),
            year=None if row["ano"] is None else int(row["ano"]),
            relator=None if row["relator"] is None else str(row["relator"]),
            natureza=natureza,
            tipo=str(row["tipo"]),
            text=str(row["texto"]),
            primary_number_raw=primary_raw,
            primary_number=primary,
            identity_block=block,
            identity_source=source,
            class_signature=signature,
            organ=organ,
        )
    return records


OCR_TRANSLATION = str.maketrans(
    {"O": "0", "o": "0", "I": "1", "i": "1", "l": "1", "S": "5", "s": "5", "G": "9", "g": "9"}
)
NUMBERISH = re.compile(r"(?<!\w)([0-9][0-9.\- /\n\t\u00a0OgGIlSso]*[0-9OgGIlSso])")


def tolerant_numeric_tokens(value: str) -> list[str]:
    tokens: list[tuple[int, int, str]] = []
    for match in NUMBERISH.finditer(value):
        normalized = re.sub(r"\D", "", match.group(1).translate(OCR_TRANSLATION))
        if normalized:
            tokens.append((len(normalized), match.start(), normalized))
    return [token for _, _, token in sorted(tokens, key=lambda item: (-item[0], item[1], item[2]))]


def primary_number_index(case_index: CaseIndex) -> dict[str, tuple[int, ...]]:
    return {
        case.numero_normalizado: case.canonical_ids
        for case in case_index.cases()
    }


def structural_number(value: str, number_index: Mapping[str, tuple[int, ...]]) -> tuple[str | None, tuple[int, ...]]:
    """Build a case-number query without consulting a gold ID.

    The longest numeric token is first normalized with the same small OCR map
    used by the diagnostic probes. Corpus keys are only used as a lookup; no
    candidate is selected by knowing the expected answer.
    """
    for token in tolerant_numeric_tokens(value):
        direct = number_index.get(token)
        if direct is not None:
            return token, direct
        contained = [
            (number, ids)
            for number, ids in number_index.items()
            if len(number) >= 4 and number in token
        ]
        if contained:
            longest = max(len(number) for number, _ in contained)
            best = [(number, ids) for number, ids in contained if len(number) == longest]
            if len(best) == 1:
                return best[0]
    return None, ()


def inferred_cnj_tribunal(number: str | None, value: str) -> str | None:
    explicit = re.search(r"\b(STF|STJ|TSE|TST|STM)\b", value, re.I)
    if explicit:
        return explicit.group(1).upper()
    if number and len(number) >= 14:
        branch = number[-7]
        return {"5": "TST", "6": "TSE", "7": "STM"}.get(branch)
    return None


def inferred_sumula_tribunal(value: str) -> str | None:
    if re.search(r"\bS[UÚ]MULA\s+VINCULANTE\b", value, re.I):
        return "STF"
    explicit = re.search(r"\b(STF|STJ|TSE|TST|STM)\b", value, re.I)
    return None if explicit is None else explicit.group(1).upper()


def structural_query(
    value: str,
    number_index: Mapping[str, tuple[int, ...]],
    records: Mapping[int, CorpusRecord],
) -> StructuralQuery:
    number, candidate_ids = structural_number(value, number_index)
    signature = class_signature(value)
    class_ids = tuple(
        candidate_id
        for candidate_id in candidate_ids
        if signature is not None and records[candidate_id].class_signature == signature
    )
    tribunal = inferred_cnj_tribunal(number, value)
    tribunal_ids = tuple(
        candidate_id
        for candidate_id in candidate_ids
        if tribunal is not None and (records[candidate_id].tribunal or "").upper() == tribunal
    )
    return StructuralQuery(number, candidate_ids, signature, class_ids, tribunal, tribunal_ids)


def is_cnj_primary(record: CorpusRecord | None) -> bool:
    return bool(record and record.primary_number_raw and CNJ_PRIMARY_PATTERN.search(record.primary_number_raw))


def oracle_family_for(
    item: GoldRow,
    span_text: str,
    records: Mapping[int, CorpusRecord],
) -> str:
    target = records.get(item.canonical_id) if item.canonical_id is not None else None
    if target and target.natureza == "dispositivo":
        return "lei_dispositivo_com_diploma" if DIPLOMA_PATTERN.search(span_text) else "lei_dispositivo_sem_diploma"
    if target and target.natureza == "sumula":
        return "sumula_numerada"
    if is_cnj_primary(target):
        return "processo_cnj"
    if target and target.natureza == "acordao" and re.search(r"\d", span_text):
        return "processo_ou_recurso_numerado"
    return surface_family(span_text, item.citation_type)


def parsed_with_data(
    parsed: ParsedCitation,
    *,
    family: str | None = None,
    data: Mapping[str, str] | None = None,
    provenance: Mapping[str, str] | None = None,
) -> ParsedCitation:
    selected_family = family or parsed.family
    return ParsedCitation(
        family=selected_family,
        tipo="lei" if selected_family.startswith("lei_") else "jurisprudencia",
        tribunal_raw=parsed.tribunal_raw,
        tribunal=parsed.tribunal,
        tribunal_source=parsed.tribunal_source,
        data=dict(parsed.data) if data is None else data,
        provenance=dict(parsed.provenance) if provenance is None else provenance,
    )


def oracle_parser_representation(
    item: GoldRow,
    span_text: str,
    family: str,
    parser: CitationParser,
    number_index: Mapping[str, tuple[int, ...]],
) -> ParsedCitation:
    """Analysis-only ideal representation; never used by production.

    Family comes from the gold target nature only in this oracle stage. Values
    are still extracted from the citation surface with a fixed, gold-free OCR
    normalization. Corpus lookup remains a separate resolver stage.
    """
    current = parser.parse(CitationCandidate(0, len(span_text), span_text, "oracle", family))
    if family == "sumula_numerada":
        match = SUMULA_TOLERANT.search(span_text)
        if match is None:
            return current
        data = dict(current.data)
        data["sumula_numero_raw"] = match.group("number")
        data["sumula_numero"] = match.group("number")
        provenance = dict(current.provenance)
        provenance["sumula_numero"] = "analysis_ocr_repair"
        return parsed_with_data(current, data=data, provenance=provenance)
    if family in {"processo_cnj", "processo_ou_recurso_numerado"}:
        tokens = tolerant_numeric_tokens(span_text)
        number = tokens[0] if tokens else None
        if number is None:
            return current
        data = dict(current.data)
        data["numero_raw"] = number
        data["numero_normalizado"] = number
        data["numero_family"] = "cnj" if family == "processo_cnj" else "case_number"
        signature = class_signature(span_text)
        if signature:
            data["classe_raw"] = signature
        provenance = dict(current.provenance)
        provenance["numero"] = "analysis_ocr_or_separator_repair"
        return parsed_with_data(current, data=data, provenance=provenance)
    return current


def correct_result(result: ResolutionResult | None, item: GoldRow) -> bool:
    return bool(result and result.status == "resolved" and result.id_canonico == item.canonical_id)


def tst_primary_guard_ids(
    query: StructuralQuery,
    records: Mapping[int, CorpusRecord],
) -> tuple[int, ...]:
    """Keep exact identity, but abstain on a TST primary-class conflict."""
    kept: list[int] = []
    for candidate_id in query.candidate_ids:
        record = records[candidate_id]
        if (record.tribunal or "").upper() != "TST":
            kept.append(candidate_id)
            continue
        if query.class_signature and record.class_signature == query.class_signature:
            kept.append(candidate_id)
    return tuple(kept)


def ideal_structural_result(
    item: GoldRow,
    span_text: str,
    records: Mapping[int, CorpusRecord],
    number_index: Mapping[str, tuple[int, ...]],
) -> tuple[ResolutionResult, StructuralQuery]:
    target = records.get(item.canonical_id) if item.canonical_id is not None else None
    query = structural_query(span_text, number_index, records)
    if target is None:
        return result_from_ids((), "ideal_structural", "acordao"), query
    if target.natureza == "acordao":
        if is_cnj_primary(target):
            ids = tst_primary_guard_ids(query, records)
            return result_from_ids(ids, "cnj_primary_class_guard", "acordao"), query
        ids = query.candidate_ids
        if len(ids) > 1 and query.class_candidate_ids:
            ids = query.class_candidate_ids
        return result_from_ids(ids, "case_number_class_filter", "acordao"), query
    # Legal identity and the two numberless sumula records deliberately remain
    # unavailable here. Sumula 211 is recovered at Stage C by Resolver V2.
    return ResolutionResult(
        "insufficient",
        None,
        (),
        None,
        "canonical_identity_not_available_in_corpus_metadata",
        "dispositivo" if target.natureza == "dispositivo" else "sumula",
    ), query


def relevant_predictions(item: GoldRow, predictions: Sequence[Prediction]) -> list[Prediction]:
    related = [
        prediction
        for prediction in predictions
        if prediction.document_id == item.document_id
        and (
            iou(item.start, item.end, prediction.start, prediction.end) > 0
            or 0 <= prediction.start - item.end <= 20
            or 0 <= item.start - prediction.end <= 20
        )
    ]
    return sorted(related, key=lambda prediction: (prediction.start, prediction.end, prediction.index))


def parsed_dict(parsed: ParsedCitation | None) -> dict[str, object] | None:
    if parsed is None:
        return None
    return {
        "family": parsed.family,
        "tipo": parsed.tipo,
        "tribunal_raw": parsed.tribunal_raw,
        "tribunal": parsed.tribunal,
        "tribunal_source": parsed.tribunal_source,
        "data": dict(parsed.data),
        "provenance": dict(parsed.provenance),
    }


def result_dict(result: ResolutionResult | None) -> dict[str, object] | None:
    return None if result is None else asdict(result)


def prediction_dict(prediction: Prediction, item: GoldRow | None = None) -> dict[str, object]:
    return {
        "index": prediction.index,
        "document": prediction.document_id,
        "span": [prediction.start, prediction.end],
        "text": prediction.text,
        "rule": prediction.rule,
        "family": prediction.family,
        "iou": None if item is None else round(iou(item.start, item.end, prediction.start, prediction.end), 6),
        "parsed": parsed_dict(prediction.parsed),
        "result": result_dict(prediction.result),
    }


def corpus_record_dict(record: CorpusRecord | None, *, snippet: int = 500) -> dict[str, object] | None:
    if record is None:
        return None
    return {
        "id_canonico": record.canonical_id,
        "documento_id": record.document_id,
        "tribunal": record.tribunal,
        "ano": record.year,
        "relator": record.relator,
        "natureza": record.natureza,
        "tipo": record.tipo,
        "primary_number_raw": record.primary_number_raw,
        "primary_number": record.primary_number,
        "identity_source": record.identity_source,
        "class_signature": record.class_signature,
        "organ": record.organ,
        "identity_block": record.identity_block,
        "text_snippet": record.text[:snippet],
    }


def result_status(result: ResolutionResult | None) -> str:
    return "missing" if result is None else result.status


def stage_and_root_audits(
    unresolved: Sequence[Evaluation],
    texts: Mapping[str, str],
    predictions: Sequence[Prediction],
    parser: CitationParser,
    resolver: CitationResolver,
    records: Mapping[int, CorpusRecord],
    number_index: Mapping[str, tuple[int, ...]],
    legal_article_index: Mapping[str, tuple[int, ...]],
) -> list[dict[str, object]]:
    audits: list[dict[str, object]] = []
    for evaluation in unresolved:
        item = evaluation.gold
        span_text = texts[item.document_id][item.start : item.end]
        target = records.get(item.canonical_id) if item.canonical_id is not None else None
        family = oracle_family_for(item, span_text, records)
        stage_b_parsed = parser.parse(CitationCandidate(0, len(span_text), span_text, "oracle_detector", family))
        stage_b_result = resolver.resolve(stage_b_parsed)
        stage_c_parsed = oracle_parser_representation(item, span_text, family, parser, number_index)
        stage_c_result = resolver.resolve(stage_c_parsed)
        stage_d_result, query = ideal_structural_result(item, span_text, records, number_index)

        stage_b_correct = correct_result(stage_b_result, item)
        stage_c_correct = correct_result(stage_c_result, item)
        stage_d_correct = correct_result(stage_d_result, item)

        if target and target.natureza == "dispositivo":
            root_cause = "legal_identity_unavailable"
            actionable = "corpus_identity"
            eventual = "verified legal source identity (article + diploma/law), then exact resolution"
        elif target and target.natureza == "sumula":
            if stage_c_correct:
                root_cause = "parser_representation_loss"
                actionable = "parser"
                eventual = "bounded OCR normalization for the sumula token and number"
            else:
                root_cause = "missing_corpus_metadata"
                actionable = "corpus_identity"
                eventual = "canonical sumula number metadata"
        elif is_cnj_primary(target):
            if item.canonical_id not in query.candidate_ids:
                root_cause = "canonical_identity_mismatch"
                actionable = "gold_issue"
                eventual = "clarify primary-identity versus source-document/body-mention semantics"
            elif len(query.candidate_ids) > 1:
                root_cause = "corpus_duplication"
                actionable = "irreducible/ambiguous"
                eventual = "corpus deduplication or an explicit version identity"
            else:
                root_cause = "cnj_resolver_policy"
                actionable = "resolver_structural"
                eventual = "primary CNJ exact resolution with a procedural-class consistency guard"
        elif stage_b_correct:
            root_cause = "detector_still_insufficient"
            actionable = "detector"
            eventual = "broader but still structural numbered-citation detection"
        elif stage_c_correct:
            root_cause = "parser_representation_loss"
            actionable = "parser"
            eventual = "bounded OCR/separator normalization after a complete span"
        elif stage_d_correct:
            root_cause = "multi_id_ambiguity"
            actionable = "resolver_structural"
            eventual = "metadata/class-aware filtering before unique selection"
        else:
            root_cause = "semantic_contextual_identity"
            actionable = "semantic_retrieval"
            eventual = "contextual candidate retrieval"

        article_match = LEGAL_ARTICLE.search(span_text)
        article = article_match.group("number") if article_match else None
        identity_candidates = query.candidate_ids
        if target and target.natureza == "dispositivo" and article:
            identity_candidates = legal_article_index.get(article, ())
        if target and target.natureza == "sumula":
            identity_candidates = stage_c_result.candidate_ids

        context_200 = texts[item.document_id][max(0, item.start - 200) : min(len(texts[item.document_id]), item.end + 200)]
        context_500 = texts[item.document_id][max(0, item.start - 500) : min(len(texts[item.document_id]), item.end + 500)]
        signals = {
            "tribunal": query.tribunal or stage_c_parsed.tribunal,
            "relator_in_candidate": bool(RELATOR_PATTERN.search(span_text)),
            "relator_in_context_200": bool(RELATOR_PATTERN.search(context_200)),
            "year_in_candidate": YEAR_PATTERN.search(span_text).group(0) if YEAR_PATTERN.search(span_text) else None,
            "year_in_context_200": YEAR_PATTERN.search(context_200).group(0) if YEAR_PATTERN.search(context_200) else None,
            "class_signature": query.class_signature,
            "organ_in_candidate": organ_from_text(span_text),
            "organ_in_context_500": organ_from_text(context_500),
            "target_metadata": None
            if target is None
            else {
                "tribunal": target.tribunal,
                "ano": target.year,
                "relator": target.relator,
                "class_signature": target.class_signature,
                "organ": target.organ,
            },
            "available_only_in_body_text": bool(
                target
                and query.number
                and query.number not in (target.primary_number or "")
                and query.number in re.sub(r"\D", "", target.text)
            ),
        }
        relevant = relevant_predictions(item, predictions)
        immediate_dependencies = [actionable]
        if evaluation.blocker in {
            "detector_miss",
            "detector_boundary_failure",
            "detector_family_mismatch",
        } and actionable != "detector":
            immediate_dependencies.insert(0, "detector_representation_or_overlap_arbitration")
        audits.append(
            {
                "documento": item.document_id,
                "nivel": item.level,
                "gold_index": item.index,
                "gold_gid": item.gid,
                "gold_family": evaluation.surface_family,
                "oracle_family": family,
                "gold_span": [item.start, item.end],
                "gold_trecho": span_text,
                "gold_id": item.canonical_id,
                "detector_candidates_relevantes": [prediction_dict(prediction, item) for prediction in relevant],
                "matched_candidate": None if evaluation.prediction is None else prediction_dict(evaluation.prediction, item),
                "matched_iou": round(evaluation.matched_iou, 6),
                "current_parsed": None if evaluation.prediction is None else parsed_dict(evaluation.prediction.parsed),
                "current_resolution": result_dict(evaluation.current_result),
                "surface_oracle_parsed": parsed_dict(evaluation.surface_oracle_parsed),
                "surface_oracle_resolution": result_dict(evaluation.surface_oracle_result),
                "oracle_parsed": parsed_dict(stage_c_parsed),
                "oracle_resolution": result_dict(stage_c_result),
                "case_index": {
                    "normalized_number": query.number,
                    "candidate_ids": list(query.candidate_ids),
                    "class_signature": query.class_signature,
                    "class_candidate_ids": list(query.class_candidate_ids),
                    "tribunal": query.tribunal,
                    "tribunal_candidate_ids": list(query.tribunal_candidate_ids),
                    "identity_candidate_ids": list(identity_candidates),
                    "identity_candidate_records": [
                        corpus_record_dict(records[candidate_id])
                        for candidate_id in identity_candidates
                        if candidate_id in records
                    ],
                },
                "first_blocker": evaluation.blocker,
                "root_cause": root_cause,
                "actionable_component": actionable,
                "immediate_dependencies": immediate_dependencies,
                "eventual_required_capability": eventual,
                "recovery_ladder": {
                    "stage_a_current": result_dict(evaluation.current_result),
                    "stage_b_oracle_detector": result_dict(stage_b_result),
                    "stage_c_oracle_parser": result_dict(stage_c_result),
                    "stage_d_ideal_structural": result_dict(stage_d_result),
                    "recovered_by_oracle_detector": stage_b_correct,
                    "recovered_by_oracle_parser": stage_c_correct and not stage_b_correct,
                    "structurally_resolvable": stage_b_correct or stage_c_correct or stage_d_correct,
                    "semantically_required": actionable == "semantic_retrieval",
                    "gold_corpus_conflict": root_cause == "canonical_identity_mismatch",
                },
                "runtime_available_evidence": {
                    "citation_text": span_text,
                    "normalized_number": query.number,
                    "class_signature": query.class_signature,
                    "tribunal": query.tribunal,
                    "context_200": context_200,
                    "context_500": context_500,
                },
                "gold_only_evidence": {
                    "target_id": item.canonical_id,
                    "target_record": corpus_record_dict(target),
                },
                "signals": signals,
            }
        )
    return sorted(audits, key=lambda audit: int(audit["gold_index"]))


def cnj_query_for_gold(
    item: GoldRow,
    span_text: str,
    parser: CitationParser,
    number_index: Mapping[str, tuple[int, ...]],
    records: Mapping[int, CorpusRecord],
) -> StructuralQuery:
    parsed = parser.parse(CitationCandidate(0, len(span_text), span_text, "cnj_oracle", "processo_cnj"))
    number = parsed.data.get("numero_normalizado")
    if number is None:
        number, _ = structural_number(span_text, number_index)
    ids = number_index.get(number, ()) if number else ()
    signature = class_signature(span_text)
    class_ids = tuple(
        candidate_id for candidate_id in ids if signature and records[candidate_id].class_signature == signature
    )
    tribunal = inferred_cnj_tribunal(number, span_text)
    tribunal_ids = tuple(
        candidate_id
        for candidate_id in ids
        if tribunal and (records[candidate_id].tribunal or "").upper() == tribunal
    )
    return StructuralQuery(number, ids, signature, class_ids, tribunal, tribunal_ids)


def strategy_outcome(
    items: Sequence[GoldRow],
    ids_by_index: Mapping[int, tuple[int, ...]],
    support_by_index: Mapping[int, bool],
) -> dict[str, int]:
    counts = Counter()
    for item in items:
        if support_by_index.get(item.index, False):
            counts["support"] += 1
        ids = ids_by_index.get(item.index, ())
        if not ids:
            counts["no_match"] += 1
        elif len(ids) > 1:
            counts["ambiguous"] += 1
        elif item.classification != "real":
            counts["false_real"] += 1
        elif ids[0] == item.canonical_id:
            counts["correct_unique"] += 1
        else:
            counts["wrong_unique"] += 1
    return {
        key: counts[key]
        for key in ("support", "correct_unique", "wrong_unique", "ambiguous", "no_match", "false_real")
    }


def body_mention_class_ids(
    query: StructuralQuery, records: Mapping[int, CorpusRecord]
) -> tuple[int, ...]:
    """Diagnostic body lookup only after a primary TST class conflict.

    Requiring an existing primary candidate avoids turning invented/no-match
    CNJs into an expensive corpus scan. The literal comes from that primary
    candidate's parsed identity, never from gold.
    """
    if len(query.candidate_ids) != 1 or not query.class_signature:
        return ()
    primary = records[query.candidate_ids[0]]
    if (primary.tribunal or "").upper() != "TST" or primary.class_signature == query.class_signature:
        return ()
    match = CNJ_PRIMARY_PATTERN.search(primary.primary_number_raw or "")
    if match is None:
        return ()
    literal = match.group(0)
    return tuple(
        record.canonical_id
        for record in records.values()
        if record.natureza == "acordao"
        and literal in record.text
        and record.class_signature == query.class_signature
    )


def cnj_query_for_prediction(
    prediction: Prediction,
    document_text: str,
    number_index: Mapping[str, tuple[int, ...]],
    records: Mapping[int, CorpusRecord],
) -> StructuralQuery:
    """Bind runtime metadata to the detected CNJ, not to an incidental window hit."""
    number = prediction.parsed.data.get("numero_normalizado")
    ids = number_index.get(number, ()) if number else ()
    prefix = document_text[max(0, prediction.start - 120) : prediction.end]
    cnj_matches = list(CNJ_PRIMARY_PATTERN.finditer(prefix))
    linked_context = prefix
    if cnj_matches:
        current_match = cnj_matches[-1]
        linked_context = prefix[max(0, current_match.start() - 80) : current_match.end()]
    signature = class_signature(linked_context)
    class_ids = tuple(
        candidate_id
        for candidate_id in ids
        if signature is not None and records[candidate_id].class_signature == signature
    )
    tribunal = inferred_cnj_tribunal(number, linked_context)
    tribunal_ids = tuple(
        candidate_id
        for candidate_id in ids
        if tribunal is not None and (records[candidate_id].tribunal or "").upper() == tribunal
    )
    return StructuralQuery(number, ids, signature, class_ids, tribunal, tribunal_ids)


def cnj_strategy_matrix(
    items: Sequence[GoldRow],
    texts: Mapping[str, str],
    parser: CitationParser,
    number_index: Mapping[str, tuple[int, ...]],
    records: Mapping[int, CorpusRecord],
) -> tuple[list[dict[str, object]], dict[int, StructuralQuery], dict[str, dict[int, tuple[int, ...]]]]:
    queries: dict[int, StructuralQuery] = {}
    for item in items:
        span_text = texts[item.document_id][item.start : item.end]
        queries[item.index] = cnj_query_for_gold(item, span_text, parser, number_index, records)

    exact = {index: query.candidate_ids for index, query in queries.items()}
    tribunal = {index: query.tribunal_candidate_ids for index, query in queries.items()}
    nature = dict(exact)  # every indexed primary identity is an acordao
    class_guard = {index: tst_primary_guard_ids(query, records) for index, query in queries.items()}
    tribunal_class = {
        index: tuple(candidate_id for candidate_id in class_guard[index] if candidate_id in query.tribunal_candidate_ids)
        for index, query in queries.items()
    }
    non_tst = {
        index: tuple(
            candidate_id for candidate_id in query.candidate_ids if (records[candidate_id].tribunal or "").upper() != "TST"
        )
        for index, query in queries.items()
    }

    # Diagnostic only: when a unique primary TST candidate conflicts in class,
    # search exact body mentions and require the mentioning record's own primary
    # class to agree. This changes identity semantics and is not promotion-ready.
    body_fallback: dict[int, tuple[int, ...]] = {}
    for index, query in queries.items():
        guarded = class_guard[index]
        if guarded or not query.number or not query.class_signature:
            body_fallback[index] = guarded
            continue
        body_fallback[index] = body_mention_class_ids(query, records)

    strategies = {
        "cnj_exact": exact,
        "tribunal_cnj": tribunal,
        "nature_cnj": nature,
        "non_tst_primary_exact": non_tst,
        "primary_class_guard": class_guard,
        "tribunal_class_cnj": tribunal_class,
        "body_mention_class_fallback_diagnostic": body_fallback,
    }
    support = {
        "cnj_exact": {index: bool(query.number) for index, query in queries.items()},
        "tribunal_cnj": {index: bool(query.number and query.tribunal) for index, query in queries.items()},
        "nature_cnj": {index: bool(query.number) for index, query in queries.items()},
        "non_tst_primary_exact": {index: bool(query.number) for index, query in queries.items()},
        "primary_class_guard": {index: bool(query.number) for index, query in queries.items()},
        "tribunal_class_cnj": {
            index: bool(query.number and query.tribunal and query.class_signature) for index, query in queries.items()
        },
        "body_mention_class_fallback_diagnostic": {
            index: bool(query.number and query.class_signature) for index, query in queries.items()
        },
    }
    matrix = []
    for name, ids in strategies.items():
        row: dict[str, object] = {"strategy": name}
        row.update(strategy_outcome(items, ids, support[name]))
        matrix.append(row)
    return matrix, queries, strategies


def transform_predictions(
    predictions: Sequence[Prediction],
    result_for: Mapping[int, ResolutionResult],
) -> list[Prediction]:
    return [replace(prediction, result=result_for.get(prediction.index, prediction.result)) for prediction in predictions]


def apply_cnj_candidate_strategy(
    strategy: str,
    predictions: Sequence[Prediction],
    texts: Mapping[str, str],
    number_index: Mapping[str, tuple[int, ...]],
    records: Mapping[int, CorpusRecord],
) -> list[Prediction]:
    replacements: dict[int, ResolutionResult] = {}
    for prediction in predictions:
        if prediction.family != "processo_cnj":
            continue
        query = cnj_query_for_prediction(
            prediction, texts[prediction.document_id], number_index, records
        )
        number = query.number
        ids = query.candidate_ids
        if strategy == "cnj_exact":
            selected = ids
        elif strategy == "non_tst_primary_exact":
            selected = tuple(cid for cid in ids if (records[cid].tribunal or "").upper() != "TST")
        elif strategy in {"primary_class_guard", "tribunal_class_cnj"}:
            selected = tst_primary_guard_ids(query, records)
            if strategy == "tribunal_class_cnj" and query.tribunal:
                selected = tuple(cid for cid in selected if (records[cid].tribunal or "").upper() == query.tribunal)
        elif strategy == "body_mention_class_fallback_diagnostic":
            selected = tst_primary_guard_ids(query, records)
            if not selected and number and query.class_signature:
                selected = body_mention_class_ids(query, records)
        else:
            raise ValueError(f"unknown CNJ strategy: {strategy}")
        replacements[prediction.index] = result_from_ids(selected, strategy, "acordao")
    return transform_predictions(predictions, replacements)


def build_legal_article_index(records: Mapping[int, CorpusRecord]) -> dict[str, tuple[int, ...]]:
    grouped: dict[str, list[int]] = defaultdict(list)
    for record in records.values():
        if record.natureza != "dispositivo":
            continue
        match = re.match(r"\s*Art[.]?\s*(\d+)", record.text, re.I)
        if match:
            grouped[match.group(1)].append(record.canonical_id)
    return {article: tuple(sorted(ids)) for article, ids in grouped.items()}


def apply_article_only(
    predictions: Sequence[Prediction],
    article_index: Mapping[str, tuple[int, ...]],
) -> list[Prediction]:
    replacements: dict[int, ResolutionResult] = {}
    for prediction in predictions:
        if not prediction.family.startswith("lei_"):
            continue
        article = prediction.parsed.data.get("artigo")
        if article:
            replacements[prediction.index] = result_from_ids(
                article_index.get(article, ()), "legal_article_only", "dispositivo"
            )
    return transform_predictions(predictions, replacements)


def apply_number_class_filter(
    predictions: Sequence[Prediction],
    records: Mapping[int, CorpusRecord],
) -> list[Prediction]:
    replacements: dict[int, ResolutionResult] = {}
    for prediction in predictions:
        if prediction.result.status != "ambiguous" or prediction.family != "processo_ou_recurso_numerado":
            continue
        signature = class_signature(prediction.text)
        selected = tuple(
            candidate_id
            for candidate_id in prediction.result.candidate_ids
            if signature and records[candidate_id].class_signature == signature
        )
        replacements[prediction.index] = result_from_ids(
            selected, "case_number_primary_class_filter", "acordao"
        )
    return transform_predictions(predictions, replacements)


def formal_fp_summary(
    gold: Sequence[GoldRow], predictions: Sequence[Prediction]
) -> dict[str, object]:
    matched = match_predictions(gold, predictions)
    matched_predictions = set(matched.values())
    formal = [prediction for prediction in predictions if prediction.index not in matched_predictions]
    return {
        "count": len(formal),
        "status": {status: sum(prediction.result.status == status for prediction in formal) for status in STATUS_ORDER[:-1]},
        "resolved": [prediction_dict(prediction) for prediction in formal if prediction.result.status == "resolved"],
    }


def resolved_formal_fp_details(
    gold: Sequence[GoldRow], predictions: Sequence[Prediction]
) -> list[dict[str, object]]:
    matched_predictions = set(match_predictions(gold, predictions).values())
    details: list[dict[str, object]] = []
    for prediction in predictions:
        if prediction.index in matched_predictions or prediction.result.status != "resolved":
            continue
        nearby = [
            item
            for item in gold
            if item.document_id == prediction.document_id
            and (
                iou(item.start, item.end, prediction.start, prediction.end) > 0
                or 0 <= prediction.start - item.end <= 20
                or 0 <= item.start - prediction.end <= 20
            )
        ]
        nearest = max(
            nearby,
            key=lambda item: (
                iou(item.start, item.end, prediction.start, prediction.end),
                -abs(item.start - prediction.start),
            ),
            default=None,
        )
        details.append(
            {
                "prediction": prediction_dict(prediction, nearest),
                "nearest_gold_index": None if nearest is None else nearest.index,
                "nearest_gold_class": None if nearest is None else nearest.classification,
                "nearest_gold_id": None if nearest is None else nearest.canonical_id,
                "resolved_matches_nearest_gold_id": bool(
                    nearest is not None
                    and nearest.classification == "real"
                    and prediction.result.id_canonico == nearest.canonical_id
                ),
                "relationship": "nested_or_overlapping" if nearest is not None and iou(
                    nearest.start, nearest.end, prediction.start, prediction.end
                ) > 0 else "adjacent_or_unrelated",
            }
        )
    return details


def status_matrix(evaluations: Sequence[Evaluation]) -> dict[str, dict[str, int]]:
    matrix: dict[str, dict[str, int]] = {}
    for classification in ("real", "inventada", "incompleta"):
        counts = Counter(
            result_status(item.current_result)
            for item in evaluations
            if item.gold.classification == classification
        )
        matrix[classification] = {status: counts[status] for status in STATUS_ORDER}
    return matrix


def normalized_person(value: str | None) -> str:
    words = compact_words(value or "")
    words = re.sub(r"\b(?:MINISTRO|MINISTRA|MIN|RELATOR|RELATORA|RELATORIA|DE)\b", " ", words)
    return " ".join(words.split())


def single_ocr_person_variants(value: str) -> tuple[str, ...]:
    """Return bounded one-character OCR variants, without fuzzy thresholds."""
    variants: set[str] = set()
    for index, character in enumerate(value):
        replacement = {"L": "I", "I": "L"}.get(character)
        if replacement is not None:
            variants.add(value[:index] + replacement + value[index + 1 :])
    return tuple(sorted(variants))


def legal_law_identity(value: str) -> tuple[str | None, str | None]:
    """Extract law number and year separately for audit evidence."""
    match = LEGAL_LAW.search(value)
    if match is None:
        return None, None
    number = re.sub(r"\D", "", match.group("number")) or None
    year = match.group("year")
    if year and len(year) == 2:
        year = ("19" if int(year) >= 30 else "20") + year
    return number, year


def incomplete_metadata_control(
    gold: Sequence[GoldRow], records: Mapping[int, CorpusRecord]
) -> dict[str, object]:
    rows = [record for record in records.values() if record.natureza == "acordao"]
    details: list[dict[str, object]] = []
    bins = Counter()
    for item in gold:
        if item.classification != "incompleta":
            continue
        court_match = COURT_PATTERN.search(item.text)
        year_match = YEAR_PATTERN.search(item.text)
        relator_match = re.search(
            r"(?:relatoria\s+(?:de|dc)|Rel[.]?\s*(?:Min[.]?|Ministro|Ministra)?)\s*(?P<name>.+)$",
            item.text,
            re.I | re.S,
        )
        if not (court_match and year_match and relator_match):
            continue
        court_words = compact_words(court_match.group(0))
        court = {
            "SUPREMO TRIBUNAL FEDERAL": "STF",
            "SUPERIOR TRIBUNAL DE JUSTICA": "STJ",
            "TRIBUNAL SUPERIOR ELEITORAL": "TSE",
            "TRIBUNAL SUPERIOR DO TRABALHO": "TST",
            "SUPERIOR TRIBUNAL MILITAR": "STM",
        }.get(court_words, court_words)
        relator = normalized_person(relator_match.group("name"))
        pool = [
            record
            for record in rows
            if (record.tribunal or "").upper() == court and str(record.year) == year_match.group(0)
        ]
        candidates = [
            record.canonical_id
            for record in pool
            if relator
            and (
                relator == normalized_person(record.relator)
                or relator in normalized_person(record.relator)
                or normalized_person(record.relator) in relator
            )
        ]
        repair_used = False
        repair_kind = None
        if not candidates and relator:
            variants = single_ocr_person_variants(relator)
            candidates = [
                record.canonical_id
                for record in pool
                if any(
                    variant == normalized_person(record.relator)
                    or variant in normalized_person(record.relator)
                    or normalized_person(record.relator) in variant
                    for variant in variants
                )
            ]
            repair_used = bool(candidates)
            repair_kind = "single L/I OCR substitution" if repair_used else None
        bucket = "0" if not candidates else "1" if len(candidates) == 1 else "2-5" if len(candidates) <= 5 else ">5"
        bins[bucket] += 1
        details.append(
            {
                "gold_index": item.index,
                "documento": item.document_id,
                "trecho": item.text,
                "tribunal": court,
                "ano": year_match.group(0),
                "relator": relator,
                "candidate_ids": candidates,
                "bucket": bucket,
                "ocr_repair_used": repair_used,
                "ocr_repair_kind": repair_kind,
            }
        )
    incomplete = [item for item in gold if item.classification == "incompleta"]
    incomplete_by_index = {item.index: item for item in incomplete}
    family_totals = Counter(surface_family(item.text, item.citation_type) for item in incomplete)
    family_triad = Counter(
        surface_family(
            incomplete_by_index[int(detail["gold_index"])].text,
            incomplete_by_index[int(detail["gold_index"])].citation_type,
        )
        for detail in details
    )
    return {
        "total_incomplete": len(incomplete),
        "cases_with_tribunal_year_relator": len(details),
        "remaining_without_triad": len(incomplete) - len(details),
        "candidate_entropy": {bucket: bins[bucket] for bucket in ("0", "1", "2-5", ">5")},
        "family_total": dict(sorted(family_totals.items())),
        "family_with_triad": dict(sorted(family_triad.items())),
        "details": details,
        "conclusion": "none of the 27 incomplete citations with tribunal + decision year + relator yields a unique candidate",
    }


def classification_audit(
    gold: Sequence[GoldRow],
    evaluations: Sequence[Evaluation],
    predictions: Sequence[Prediction],
) -> dict[str, object]:
    matrix = status_matrix(evaluations)
    status_totals: dict[str, Counter[str]] = {
        status: Counter(
            item.gold.classification
            for item in evaluations
            if result_status(item.current_result) == status
        )
        for status in STATUS_ORDER
    }
    majority_correct = sum(max(counts.values(), default=0) for counts in status_totals.values())
    detected_majority_correct = sum(
        max(status_totals[status].values(), default=0) for status in STATUS_ORDER if status != "missing"
    )
    simple_mapping = {
        "resolved": "real",
        "no_match": "inventada",
        "ambiguous": "incompleta",
        "insufficient": "incompleta",
    }
    detected_evaluations = [item for item in evaluations if item.current_result is not None]
    simple_correct = sum(
        simple_mapping[item.current_result.status] == item.gold.classification
        for item in detected_evaluations
    )
    matched_prediction_ids = set(match_predictions(gold, predictions).values())
    emitted_by_class = Counter(simple_mapping[prediction.result.status] for prediction in predictions)
    correct_by_class = Counter(
        simple_mapping[item.current_result.status]
        for item in evaluations
        if item.current_result is not None
        and simple_mapping[item.current_result.status] == item.gold.classification
    )
    general_incomplete = [
        item
        for item in evaluations
        if item.current_result is not None
        and item.current_result.status == "insufficient"
        and item.prediction is not None
        and item.prediction.family == "jurisprudencia_referencia_geral"
    ]
    general_fp = [
        prediction
        for prediction in predictions
        if prediction.family == "jurisprudencia_referencia_geral" and prediction.index not in matched_prediction_ids
    ]
    return {
        "status_by_gold": matrix,
        "detected_by_gold": {
            classification: sum(
                item.gold.classification == classification and item.current_result is not None
                for item in evaluations
            )
            for classification in ("real", "inventada", "incompleta")
        },
        "gold_status_oracle_bound_not_runtime_achievable": {
            "correct": majority_correct,
            "total": len(gold),
            "accuracy": round(majority_correct / len(gold), 6),
            "majority_by_status": {
                status: (counts.most_common(1)[0][0] if counts else None)
                for status, counts in status_totals.items()
            },
            "warning": "includes missing gold spans and chooses each status label with gold; this is a separability diagnostic, not an executable classifier",
        },
        "detected_gold_status_oracle_bound": {
            "correct": detected_majority_correct,
            "total": len(detected_evaluations),
            "accuracy": round(detected_majority_correct / len(detected_evaluations), 6),
            "warning": "majority labels are fitted with gold and exclude the 84 formal detector false positives",
        },
        "simple_detection_gated_mapping": {
            "mapping": simple_mapping,
            "correct": simple_correct,
            "matched_gold": len(detected_evaluations),
            "matched_accuracy": round(simple_correct / len(detected_evaluations), 6),
            "gold_total": len(gold),
            "gold_coverage_accuracy": round(simple_correct / len(gold), 6),
            "detector_predictions_emitted": len(predictions),
            "exact_span_and_class_precision": round(simple_correct / len(predictions), 6),
            "formal_false_positive_outputs": len(predictions) - len(matched_prediction_ids),
            "predicted_by_class": dict(sorted(emitted_by_class.items())),
            "correct_by_predicted_class": dict(sorted(correct_by_class.items())),
            "precision_by_predicted_class": {
                classification: round(correct_by_class[classification] / emitted_by_class[classification], 6)
                for classification in ("real", "inventada", "incompleta")
            },
            "warning": "missing gold spans produce no runtime output; mapping missing to incompleta would incorrectly credit 53 false negatives",
        },
        "status_composition": {status: dict(sorted(counts.items())) for status, counts in status_totals.items()},
        "classification_only_opportunity": {
            "rule": "jurisprudencia_referencia_geral + insufficient -> incompleta",
            "observed_correct": sum(item.gold.classification == "incompleta" for item in general_incomplete),
            "support": len(general_incomplete),
            "counterexamples": sum(item.gold.classification != "incompleta" for item in general_incomplete),
            "formal_fp_support": len(general_fp),
            "unit": "class labels, not canonical IDs",
            "verdict": "bounded selective opportunity; not readiness for a complete three-way classifier",
        },
        "readiness": "NOT_READY",
        "reason": "missing and insufficient mix all three gold classes; no_match also contains three real citations",
    }


def invented_gold_audit(
    evaluations: Sequence[Evaluation], records: Mapping[int, CorpusRecord]
) -> dict[str, object]:
    invented = [item for item in evaluations if item.gold.classification == "inventada"]
    current_status = Counter(result_status(item.current_result) for item in invented)
    oracle_status = Counter(item.surface_oracle_result.status for item in invented)
    transitions: dict[str, Counter[str]] = defaultdict(Counter)
    for item in invented:
        transitions[result_status(item.current_result)][item.surface_oracle_result.status] += 1

    stf_sumula_ids = tuple(
        sorted(
            record.canonical_id
            for record in records.values()
            if record.natureza == "sumula" and (record.tribunal or "").upper() == "STF"
        )
    )
    invented_vinculante = [
        item for item in invented if re.search(r"\bS[UÚ]MULA\s+VINCULANTE\b", item.gold.text, re.I)
    ]
    tribunal_only_false_real = len(invented_vinculante) if len(stf_sumula_ids) == 1 else 0
    return {
        "total": len(invented),
        "current_status": {status: current_status[status] for status in STATUS_ORDER},
        "first_blockers": dict(sorted(Counter(item.blocker for item in invented).items())),
        "surface_families": dict(sorted(Counter(item.surface_family for item in invented).items())),
        "family_by_current_status": {
            status: dict(
                sorted(
                    Counter(
                        item.surface_family
                        for item in invented
                        if result_status(item.current_result) == status
                    ).items()
                )
            )
            for status in STATUS_ORDER
        },
        "surface_oracle_current_parser_resolver": {
            "status": {status: oracle_status[status] for status in STATUS_ORDER[:-1]},
            "runtime_to_oracle": {
                current: dict(sorted(next_status.items()))
                for current, next_status in sorted(transitions.items())
            },
            "structurally_sufficient_no_match": oracle_status["no_match"],
            "unsupported_insufficient": oracle_status["insufficient"],
            "historical_oracle_v2_reference": {"no_match": 38, "insufficient": 26},
            "historical_delta_requires_separate_harness_reconciliation": {
                "no_match": oracle_status["no_match"] - 38,
                "insufficient": oracle_status["insufficient"] - 26,
            },
            "note": "the current full-span surface-family oracle is 41/23, not the frozen historical 38/26; it is reported separately rather than silently substituted",
        },
        "structural_identity_available_now": {
            "no_match": current_status["no_match"],
            "interpretation": "these citations already form a sufficient exact query whose key is absent from the corpus",
        },
        "sumula_tribunal_only_counterfactual": {
            "runtime_rule": "Súmula Vinculante implies STF, then select a tribunal singleton",
            "stf_sumula_candidate_ids": list(stf_sumula_ids),
            "invented_vinculante_cases": [
                {"gold_index": item.gold.index, "documento": item.gold.document_id, "trecho": item.gold.text}
                for item in invented_vinculante
            ],
            "false_real_if_unique": tribunal_only_false_real,
            "verdict": "unsafe; preserve number identity and do not fall back to tribunal-only selection",
        },
        "cases": [
            {
                "gold_index": item.gold.index,
                "documento": item.gold.document_id,
                "nivel": item.gold.level,
                "trecho": item.gold.text,
                "surface_family": item.surface_family,
                "first_blocker": item.blocker,
                "current_status": result_status(item.current_result),
                "current_parsed": None if item.prediction is None else parsed_dict(item.prediction.parsed),
                "current_resolution": result_dict(item.current_result),
                "surface_oracle_parsed": parsed_dict(item.surface_oracle_parsed),
                "surface_oracle_resolution": result_dict(item.surface_oracle_result),
            }
            for item in invented
        ],
        "classification_implication": "no_match is strong invented evidence, while missing and unsupported insufficient cannot safely distinguish invented from real/incomplete",
    }


def detector_metrics(
    gold: Sequence[GoldRow], predictions: Sequence[Prediction], evaluations: Sequence[Evaluation]
) -> dict[str, object]:
    matches = match_predictions(gold, predictions)
    by_prediction = {prediction.index: prediction for prediction in predictions}
    tp = len(matches)
    fp = len(predictions) - tp
    fn = len(gold) - tp
    exact = sum(
        item.start == by_prediction[prediction_index].start
        and item.end == by_prediction[prediction_index].end
        for item in gold
        for prediction_index in [matches.get(item.index)]
        if prediction_index is not None
    )
    precision = tp / len(predictions)
    recall = tp / len(gold)
    return {
        "predictions": len(predictions),
        "TP": tp,
        "FP": fp,
        "FN": fn,
        "precision": round(precision, 6),
        "recall": round(recall, 6),
        "F1": round(2 * precision * recall / (precision + recall), 6),
        "exact": exact,
        "identity_preservation": sum(
            evaluation.prediction is not None and evaluation.parser_equivalent
            for evaluation in evaluations
        ),
    }


def aggregate_counts(items: Sequence[dict[str, object]], field: str) -> list[dict[str, object]]:
    counts = Counter(str(item[field]) for item in items)
    total = len(items)
    return [
        {"name": name, "cases": count, "percent": round(100 * count / total, 1)}
        for name, count in sorted(counts.items(), key=lambda pair: (-pair[1], pair[0]))
    ]


def audit_cnj_current_strategies(
    gold: Sequence[GoldRow],
    texts: Mapping[str, str],
    predictions: Sequence[Prediction],
    resolver: CitationResolver,
    number_index: Mapping[str, tuple[int, ...]],
    records: Mapping[int, CorpusRecord],
    baseline_evaluations: Sequence[Evaluation],
) -> dict[str, object]:
    baseline_safety = safe_metrics(baseline_evaluations)
    rows: list[dict[str, object]] = []
    evaluations_by_strategy: dict[str, list[Evaluation]] = {}
    formal_details_by_strategy: dict[str, list[dict[str, object]]] = {}
    for strategy in (
        "cnj_exact",
        "non_tst_primary_exact",
        "primary_class_guard",
        "tribunal_class_cnj",
        "body_mention_class_fallback_diagnostic",
    ):
        transformed = apply_cnj_candidate_strategy(strategy, predictions, texts, number_index, records)
        evaluations = evaluate(gold, texts, transformed, resolver)
        evaluations_by_strategy[strategy] = evaluations
        safety = safe_metrics(evaluations)
        formal = formal_fp_summary(gold, transformed)
        formal_details_by_strategy[strategy] = resolved_formal_fp_details(gold, transformed)
        rows.append(
            {
                "strategy": strategy,
                "correct_ids": safety["correct"],
                "gain": safety["correct"] - baseline_safety["correct"],
                "wrong_unique_real": safety["wrong_unique_real"],
                "false_real_inventada": safety["false_real_inventada"],
                "false_real_incompleta": safety["false_real_incompleta"],
                "formal_fp_resolved": formal["status"]["resolved"],
                "promotion_candidate": strategy == "primary_class_guard",
                "promotion_ready": False,
                "promotion_blocker": (
                    "requires deterministic overlap/duplicate arbitration for resolved formal false positives"
                    if strategy == "primary_class_guard"
                    else None
                ),
            }
        )
    return {
        "matrix": rows,
        "evaluations": evaluations_by_strategy,
        "resolved_formal_fp_details": formal_details_by_strategy,
    }


def gold_cnj_case_record(
    item: GoldRow,
    evaluation: Evaluation,
    query: StructuralQuery,
    records: Mapping[int, CorpusRecord],
) -> dict[str, object]:
    target = records.get(item.canonical_id) if item.canonical_id is not None else None
    return {
        "gold_index": item.index,
        "documento": item.document_id,
        "nivel": item.level,
        "classification": item.classification,
        "trecho": item.text,
        "gold_id": item.canonical_id,
        "detector_status": evaluation.blocker,
        "formally_matched": evaluation.prediction is not None,
        "detector_candidate": None if evaluation.prediction is None else prediction_dict(evaluation.prediction, item),
        "normalized_cnj": query.number,
        "tribunal": query.tribunal,
        "classe": query.class_signature,
        "candidate_ids": list(query.candidate_ids),
        "gold_in_candidates": item.canonical_id in query.candidate_ids if item.canonical_id is not None else False,
        "unique": len(query.candidate_ids) == 1,
        "target_primary_identity": None if target is None else target.primary_number_raw,
        "target_record": corpus_record_dict(target, snippet=250),
        "candidate_records": [corpus_record_dict(records[candidate_id], snippet=250) for candidate_id in query.candidate_ids],
    }


def build_audit() -> tuple[dict[str, object], dict[str, float]]:
    started = time.perf_counter()
    if not DATASET_DIR.is_dir() or get_database_path().resolve() != (DATASET_DIR / "desafio1_bracis.db").resolve():
        raise BaselineError("dataset invalido: use somente material_desafio_jusbrasil_bracis/")
    texts = load_texts()
    gold = load_gold()
    if (len(texts), len(gold)) != (26, 225):
        raise BaselineError(f"dataset divergente: documentos={len(texts)}, citacoes={len(gold)}")
    if Counter(item.classification for item in gold) != EXPECTED_CLASSES:
        raise BaselineError("distribuicao de classes divergente")
    if Counter(item.level for item in gold) != EXPECTED_LEVELS:
        raise BaselineError("distribuicao N1/N2 divergente")
    for item in gold:
        actual = texts[item.document_id][item.start : item.end]
        if actual != item.text:
            raise BaselineError(f"span gold divergente no index {item.index}")
    database_hash = file_sha256(get_database_path())
    if database_hash != EXPECTED_DB_HASH:
        raise BaselineError(f"SHA-256 divergente: {database_hash}")

    with connect_database(get_database_path(), read_only=True) as connection:
        integrity = str(connection.execute("PRAGMA integrity_check").fetchone()[0])
        if integrity != "ok":
            raise BaselineError(f"PRAGMA integrity_check divergente: {integrity}")
        database_counts = {
            str(row["natureza"]): int(row["count"])
            for row in connection.execute(
                "SELECT natureza, COUNT(*) AS count FROM documentos GROUP BY natureza ORDER BY natureza"
            ).fetchall()
        }
        if database_counts != {"acordao": 998, "dispositivo": 13, "sumula": 5}:
            raise BaselineError(f"contagens do banco divergentes: {database_counts}")
        database_columns = tuple(
            str(row["name"])
            for row in connection.execute("PRAGMA table_info(documentos)").fetchall()
        )
        structured_legal_identity_fields = tuple(
            field
            for field in ("diploma", "law_number", "lei_numero", "source_diploma", "source_year")
            if field in database_columns
        )
        case_index = build_case_index(connection)
        if case_index.stats != {
            "records_total": 998,
            "case_keys": 912,
            "single_id_cases": 833,
            "multi_id_cases": 79,
            "max_ids_per_case": 4,
        }:
            raise BaselineError(f"CaseIndex divergente: {case_index.stats}")
        resolver = CitationResolver(case_index=case_index, connection=connection)
        records = corpus_records(connection)
        parser = CitationParser()
        numbers = primary_number_index(case_index)
        article_index = build_legal_article_index(records)

        predictions, pipeline_seconds = run_pipeline(texts, resolver)
        evaluations = evaluate(gold, texts, predictions, resolver)
        detector = detector_metrics(gold, predictions, evaluations)
        observed_detector = tuple(int(detector[key]) for key in ("predictions", "TP", "FP", "FN", "exact"))
        if observed_detector != EXPECTED_DETECTOR:
            raise BaselineError(f"Detector V4 divergente: {observed_detector}")
        baseline_safety = safe_metrics(evaluations)
        if baseline_safety != {
            "correct": 41,
            "wrong_unique_real": 0,
            "false_real_inventada": 0,
            "false_real_incompleta": 0,
        }:
            raise BaselineError(f"E2E/safety divergente: {baseline_safety}")
        baseline_formal_fp = formal_fp_summary(gold, predictions)
        if baseline_formal_fp["count"] != 84 or baseline_formal_fp["status"] != {
            "resolved": 0,
            "no_match": 0,
            "ambiguous": 0,
            "insufficient": 84,
        }:
            raise BaselineError(f"downstream dos FPs divergente: {baseline_formal_fp}")

        real_evaluations = [item for item in evaluations if item.gold.classification == "real"]
        unresolved_evaluations = [item for item in real_evaluations if item.blocker != "resolved_correct"]
        if len(unresolved_evaluations) != 55:
            raise BaselineError(f"real unresolved divergente: {len(unresolved_evaluations)}")

        unresolved = stage_and_root_audits(
            unresolved_evaluations,
            texts,
            predictions,
            parser,
            resolver,
            records,
            numbers,
            article_index,
        )
        first_blockers = aggregate_counts(unresolved, "first_blocker")
        roots = aggregate_counts(unresolved, "root_cause")
        actionable = aggregate_counts(unresolved, "actionable_component")

        blocker_counter = Counter(str(item["first_blocker"]) for item in unresolved)
        expected_blockers = Counter(
            {
                "resolver_insufficient": 21,
                "detector_miss": 14,
                "detector_boundary_failure": 12,
                "detector_family_mismatch": 5,
                "resolver_no_match": 2,
                "resolver_ambiguous": 1,
            }
        )
        if blocker_counter != expected_blockers:
            raise BaselineError(f"first blockers divergentes: {blocker_counter}")

        stage_b_new = sum(bool(item["recovery_ladder"]["recovered_by_oracle_detector"]) for item in unresolved)
        stage_c_new = sum(bool(item["recovery_ladder"]["recovered_by_oracle_parser"]) for item in unresolved)
        stage_d_new = sum(
            correct_result(
                ResolutionResult(**item["recovery_ladder"]["stage_d_ideal_structural"]),
                next(evaluation.gold for evaluation in unresolved_evaluations if evaluation.gold.index == item["gold_index"]),
            )
            and not bool(item["recovery_ladder"]["recovered_by_oracle_detector"])
            and not bool(item["recovery_ladder"]["recovered_by_oracle_parser"])
            for item in unresolved
            if item["recovery_ladder"]["stage_d_ideal_structural"] is not None
        )
        # Compatibility diagnostic only: compare the frozen surface-family
        # harness with the root-aware oracle used by this audit.
        legacy_detector_indexes = {
            item.gold.index
            for item in unresolved_evaluations
            if item.blocker.startswith("detector_")
            and correct_result(item.surface_oracle_result, item.gold)
        }
        root_aware_detector_indexes = {
            int(item["gold_index"])
            for item in unresolved
            if bool(item["recovery_ladder"]["recovered_by_oracle_detector"])
        }
        legacy_detector_upper = len(legacy_detector_indexes)

        strict_cnj_items = [
            item
            for item in gold
            if surface_family(texts[item.document_id][item.start : item.end], item.citation_type) == "processo_cnj"
        ]
        broad_cnj_items = [
            item
            for item in gold
            if (
                item.canonical_id is not None and is_cnj_primary(records.get(item.canonical_id))
            )
            or item in strict_cnj_items
        ]
        cnj_matrix, cnj_queries, cnj_strategy_ids = cnj_strategy_matrix(
            strict_cnj_items, texts, parser, numbers, records
        )
        broad_cnj_matrix, broad_cnj_queries, broad_cnj_strategy_ids = cnj_strategy_matrix(
            broad_cnj_items, texts, parser, numbers, records
        )
        current_cnj = audit_cnj_current_strategies(
            gold, texts, predictions, resolver, numbers, records, evaluations
        )
        by_gold_eval = {item.gold.index: item for item in evaluations}
        cnj_cases = [
            gold_cnj_case_record(item, by_gold_eval[item.index], cnj_queries[item.index], records)
            for item in strict_cnj_items
        ]
        broad_cnj_cases = [
            gold_cnj_case_record(item, by_gold_eval[item.index], broad_cnj_queries[item.index], records)
            for item in broad_cnj_items
        ]

        exact_wrong = [
            item
            for item in strict_cnj_items
            if len(cnj_strategy_ids["cnj_exact"][item.index]) == 1
            and item.classification == "real"
            and cnj_strategy_ids["cnj_exact"][item.index][0] != item.canonical_id
        ]
        if len(exact_wrong) != 1:
            raise BaselineError(f"conflito CNJ unique divergente: {len(exact_wrong)}")
        conflict_item = exact_wrong[0]
        conflict_query = cnj_queries[conflict_item.index]
        conflict_candidate_id = conflict_query.candidate_ids[0]
        conflict_target = records[conflict_item.canonical_id]
        conflict_candidate = records[conflict_candidate_id]
        number_surface = CNJ_PRIMARY_PATTERN.search(conflict_candidate.primary_number_raw or "")
        cited_form = number_surface.group(0) if number_surface else ""
        mention_offset = conflict_target.text.find(cited_form)
        conflict_without_cnj = CNJ_PRIMARY_PATTERN.sub(" ", conflict_item.text)
        origin_year_match = re.search(
            r"\d{1,7}\s*-\s*\d{2}\s*[. ]\s*(?P<year>\d{4})\s*[. ]\s*\d\s*[. ]\s*\d{2}\s*[. ]\s*\d{4}",
            conflict_item.text,
        )
        conflict = {
            "gold_index": conflict_item.index,
            "documento": conflict_item.document_id,
            "trecho": conflict_item.text,
            "gold_id": conflict_item.canonical_id,
            "primary_exact_id": conflict_candidate_id,
            "gold_record": corpus_record_dict(conflict_target),
            "primary_exact_record": corpus_record_dict(conflict_candidate),
            "gold_primary_differs": conflict_target.primary_number != conflict_query.number,
            "gold_body_mention_offset_zero_based": mention_offset,
            "gold_contains_body_mention": mention_offset >= 0,
            "gold_body_mention_context_300": (
                None
                if mention_offset < 0
                else conflict_target.text[
                    max(0, mention_offset - 300) : min(
                        len(conflict_target.text), mention_offset + len(cited_form) + 300
                    )
                ]
            ),
            "citation_class": conflict_query.class_signature,
            "primary_exact_class": conflict_candidate.class_signature,
            "runtime_distinguishing_evidence": {
                "class_conflict": conflict_query.class_signature != conflict_candidate.class_signature,
                "cnj_origin_year_present": None if origin_year_match is None else origin_year_match.group("year"),
                "decision_year_explicit_outside_cnj": (
                    YEAR_PATTERN.search(conflict_without_cnj).group(0)
                    if YEAR_PATTERN.search(conflict_without_cnj)
                    else None
                ),
                "relator_in_citation": bool(RELATOR_PATTERN.search(conflict_item.text)),
                "a_record_matches_both_number_and_class": bool(conflict_query.class_candidate_ids),
            },
            "conclusion": "gold points to a source document that mentions the CNJ; primary structural identity points elsewhere",
            "avoidable_wrong_unique": True,
            "safe_action": "abstain when a TST primary class conflicts with the cited procedural prefix",
        }

        cnj_ambiguous = [
            case
            for case in cnj_cases
            if case["classification"] == "real" and len(case["candidate_ids"]) > 1
        ]
        for case in cnj_ambiguous:
            candidate_records = [records[candidate_id] for candidate_id in case["candidate_ids"]]
            case["metadata_equal"] = all(
                (
                    record.tribunal,
                    record.year,
                    normalized_person(record.relator),
                    record.class_signature,
                    record.organ,
                )
                == (
                    candidate_records[0].tribunal,
                    candidate_records[0].year,
                    normalized_person(candidate_records[0].relator),
                    candidate_records[0].class_signature,
                    candidate_records[0].organ,
                )
                for record in candidate_records[1:]
            )
            word_sets = [set(re.findall(r"\w+", record.text.lower())) for record in candidate_records]
            case["word_jaccard"] = round(
                len(word_sets[0] & word_sets[1]) / len(word_sets[0] | word_sets[1]), 6
            )
            case["unique_selection_possible"] = False

        legal_real = [
            evaluation
            for evaluation in unresolved_evaluations
            if records[evaluation.gold.canonical_id].natureza == "dispositivo"
        ]
        article_predictions = apply_article_only(predictions, article_index)
        article_evaluations = evaluate(gold, texts, article_predictions, resolver)
        article_safety = safe_metrics(article_evaluations)
        article_formal = formal_fp_summary(gold, article_predictions)
        article_formal_details = resolved_formal_fp_details(gold, article_predictions)
        article_formal_near_invented = sum(
            detail["nearest_gold_class"] == "inventada" for detail in article_formal_details
        )
        legal_cases: list[dict[str, object]] = []
        for evaluation in legal_real:
            item = evaluation.gold
            span_text = texts[item.document_id][item.start : item.end]
            oracle_family = oracle_family_for(item, span_text, records)
            parsed = parser.parse(CitationCandidate(0, len(span_text), span_text, "legal_oracle", oracle_family))
            article = parsed.data.get("artigo")
            law_number, law_year = legal_law_identity(span_text)
            article_candidates = article_index.get(article or "", ())
            legal_cases.append(
                {
                    "gold_index": item.index,
                    "documento": item.document_id,
                    "trecho": span_text,
                    "gold_id": item.canonical_id,
                    "detector_status": evaluation.blocker,
                    "current_candidate": (
                        None if evaluation.prediction is None else prediction_dict(evaluation.prediction, item)
                    ),
                    "current_parsed": (
                        None if evaluation.prediction is None else parsed_dict(evaluation.prediction.parsed)
                    ),
                    "current_resolution": result_dict(evaluation.current_result),
                    "oracle_full_span_parser": parsed_dict(parsed),
                    "artigo": article,
                    "diploma": parsed.data.get("diploma_normalizado"),
                    "law_number": law_number,
                    "year": law_year,
                    "parser_law_number_raw_behavior": parsed.data.get("law_number"),
                    "article_only_candidate_ids": list(article_candidates),
                    "article_only_unique_correct": article_candidates == (item.canonical_id,),
                    "verified_diploma_candidate_ids": [],
                }
            )
        legal_oracle_correct = sum(bool(case["article_only_unique_correct"]) for case in legal_cases)
        legal_e2e_gain = article_safety["correct"] - baseline_safety["correct"]
        legal_article_safe = all(
            article_safety[key] == 0
            for key in ("wrong_unique_real", "false_real_inventada", "false_real_incompleta")
        )
        unsafe_legal = [
            item
            for item in article_evaluations
            if item.gold.classification != "real"
            and item.current_result is not None
            and item.current_result.status == "resolved"
        ]

        ambiguity_predictions = apply_number_class_filter(predictions, records)
        ambiguity_evaluations = evaluate(gold, texts, ambiguity_predictions, resolver)
        ambiguity_safety = safe_metrics(ambiguity_evaluations)
        current_ambiguous = [
            item for item in evaluations if item.gold.classification == "real" and result_status(item.current_result) == "ambiguous"
        ]
        all_ambiguous: list[dict[str, object]] = []
        for evaluation in current_ambiguous:
            all_ambiguous.append(
                {
                    "gold_index": evaluation.gold.index,
                    "documento": evaluation.gold.document_id,
                    "trecho": evaluation.gold.text,
                    "gold_id": evaluation.gold.canonical_id,
                    "candidate_ids": list(evaluation.current_result.candidate_ids),
                    "candidate_records": [corpus_record_dict(records[candidate_id]) for candidate_id in evaluation.current_result.candidate_ids],
                    "runtime_class": class_signature(evaluation.prediction.text),
                    "unique_selection_possible": any(
                        changed.gold.index == evaluation.gold.index and changed.blocker == "resolved_correct"
                        for changed in ambiguity_evaluations
                    ),
                    "mechanism": "primary procedural class filter",
                }
            )
        all_ambiguous.extend(cnj_ambiguous)

        real_no_match = [
            {
                "gold_index": item.gold.index,
                "documento": item.gold.document_id,
                "trecho": item.gold.text,
                "gold_id": item.gold.canonical_id,
                "first_blocker": item.blocker,
                "current_result": result_dict(item.current_result),
                "root_cause": next(audit["root_cause"] for audit in unresolved if audit["gold_index"] == item.gold.index),
                "reason": (
                    "canonical sumula number absent from corpus metadata"
                    if records[item.gold.canonical_id].natureza == "sumula"
                    else "terminal OCR character O was excluded/corrupted before exact lookup"
                ),
            }
            for item in evaluations
            if item.gold.classification == "real" and result_status(item.current_result) == "no_match"
        ]
        for case in real_no_match:
            target = records[int(case["gold_id"])]
            if target.natureza != "sumula":
                continue
            tribunal = inferred_sumula_tribunal(str(case["trecho"]))
            tribunal_ids = tuple(
                sorted(
                    record.canonical_id
                    for record in records.values()
                    if record.natureza == "sumula"
                    and tribunal is not None
                    and (record.tribunal or "").upper() == tribunal
                )
            )
            case["tribunal_only_counterfactual"] = {
                "inferred_tribunal": tribunal,
                "candidate_ids": list(tribunal_ids),
                "would_be_unique": len(tribunal_ids) == 1,
                "safety_warning": "invented Súmula Vinculante controls show that tribunal-only singleton selection creates false-real",
            }

        semantic_cases = [audit for audit in unresolved if audit["actionable_component"] == "semantic_retrieval"]
        noisy_structural = [
            audit
            for audit in unresolved
            if audit["first_blocker"] == "detector_miss"
            or (
                audit["root_cause"] == "parser_representation_loss"
                and audit["first_blocker"] == "resolver_no_match"
            )
        ]
        semantic_readiness = Counter(
            "HIGH"
            if sum(
                bool(audit["signals"].get(key))
                for key in ("tribunal", "relator_in_context_200", "year_in_context_200", "class_signature")
            )
            >= 3
            else "MEDIUM"
            if sum(
                bool(audit["signals"].get(key))
                for key in ("tribunal", "relator_in_context_200", "year_in_context_200", "class_signature")
            )
            >= 1
            else "LOW"
            for audit in semantic_cases
        )
        noisy_unique_correct = sum(
            len(audit["case_index"]["identity_candidate_ids"]) == 1
            and audit["case_index"]["identity_candidate_ids"][0] == audit["gold_id"]
            for audit in noisy_structural
        )

        classification = classification_audit(gold, evaluations, predictions)
        invented_audit = invented_gold_audit(evaluations, records)
        incomplete_control = incomplete_metadata_control(gold, records)

        entropy = Counter()
        for audit in unresolved:
            candidate_count = len(audit["case_index"]["identity_candidate_ids"])
            bucket = "0" if candidate_count == 0 else "1" if candidate_count == 1 else "2-5" if candidate_count <= 5 else ">5"
            entropy[bucket] += 1

        detector_blocked = [audit for audit in unresolved if str(audit["first_blocker"]).startswith("detector_")]
        detector_residual = {
            "first_blocker_cases": len(detector_blocked),
            "detector_only_immediate_upper_bound_post_v4": stage_b_new,
            "legacy_surface_taxonomy_upper_bound": legacy_detector_upper,
            "root_aware_gold_indexes": sorted(root_aware_detector_indexes),
            "legacy_surface_gold_indexes": sorted(legacy_detector_indexes),
            "legacy_only_gold_indexes": sorted(legacy_detector_indexes - root_aware_detector_indexes),
            "root_aware_only_gold_indexes": sorted(root_aware_detector_indexes - legacy_detector_indexes),
            "legacy_warning": "aggregate comparison can conceal different cases; use the root-aware set because family is assigned from target nature only within the oracle analysis",
            "issue_classes": {
                "generalizable_structural": stage_b_new,
                "isolated_fragile": sum(
                    audit["root_cause"] == "parser_representation_loss" for audit in detector_blocked
                ),
                "semantic_context_dependent": sum(
                    audit["root_cause"] == "semantic_contextual_identity" for audit in detector_blocked
                ),
                "family_policy_downstream": sum(
                    audit["root_cause"]
                    not in {"detector_still_insufficient", "parser_representation_loss", "semantic_contextual_identity"}
                    for audit in detector_blocked
                ),
            },
            "saturation": "nearing saturation",
            "freeze_supported": True,
            "reason": f"only {stage_b_new} IDs are immediately recoverable by a root-aware detector oracle; most visible detector errors remain blocked downstream",
        }
        parser_audit = {
            "parser_only_immediate_upper_bound": sum(
                audit["first_blocker"] == "parser_gap"
                and audit["recovery_ladder"]["recovered_by_oracle_parser"]
                for audit in unresolved
            ),
            "joint_oracle_detector_parser_newly_recovered": stage_c_new,
            "cases": [
                audit["gold_index"] for audit in unresolved if audit["recovery_ladder"]["recovered_by_oracle_parser"]
            ],
            "saturation": "saturated for current candidate spans",
        }

        classification_matrix = classification["status_by_gold"]
        recovery_ladder = [
            {"stage": "current", "newly_recovered": 41, "cumulative": 41},
            {"stage": "oracle_detector", "newly_recovered": stage_b_new, "cumulative": 41 + stage_b_new},
            {
                "stage": "oracle_parser",
                "newly_recovered": stage_c_new,
                "cumulative": 41 + stage_b_new + stage_c_new,
            },
            {
                "stage": "ideal_structural",
                "newly_recovered": stage_d_new,
                "cumulative": 41 + stage_b_new + stage_c_new + stage_d_new,
            },
        ]

        root_matrix = [
            {
                "root_cause": row["name"],
                "real_cases": row["cases"],
                "immediately_actionable": row["name"]
                in {"cnj_resolver_policy", "detector_still_insufficient", "parser_representation_loss", "multi_id_ambiguity"},
                "likely_component": next(
                    audit["actionable_component"]
                    for audit in unresolved
                    if audit["root_cause"] == row["name"]
                ),
            }
            for row in roots
        ]

        cnj_safe_row = next(
            row for row in current_cnj["matrix"] if row["strategy"] == "primary_class_guard"
        )
        cnj_safe_gain = int(cnj_safe_row["gain"])
        cnj_plausible_gain = sum(
            audit["root_cause"] == "cnj_resolver_policy" for audit in unresolved
        )
        broad_cnj_remaining_real = sum(
            item.classification == "real" and by_gold_eval[item.index].blocker != "resolved_correct"
            for item in broad_cnj_items
        )

        roadmap = [
            {
                "path": "CNJ / canonical identity",
                "cases_affected": broad_cnj_remaining_real,
                "safe_known_gain": cnj_safe_gain,
                "plausible_gain": cnj_plausible_gain,
                "risk": "LOW-MEDIUM: class alias coverage and overlapping candidates; abstention preserves safety",
                "complexity": "MEDIUM",
                "generalization": "HIGH for exact primary identity; MEDIUM for procedural-class aliases on degraded surfaces",
                "gold_dependence": "none in candidate construction; gold used only for evaluation",
                "thresholds": "none",
                "evidence": "primary_class_guard gives +14 IDs with 0/0/0; broad oracle has 27 gold-consistent unique unresolved CNJs",
            },
            {
                "path": "Ambiguity resolution",
                "cases_affected": 3,
                "safe_known_gain": 1,
                "plausible_gain": 1,
                "risk": "LOW on the observed class-distinguishable STJ case; two TST duplicate pairs must abstain",
                "complexity": "LOW-MEDIUM",
                "generalization": "MEDIUM: deterministic class filtering generalizes, but observed safe gain is one case",
                "gold_dependence": "none in filtering; gold used only for evaluation",
                "thresholds": "none",
                "evidence": "number + primary class resolves the only current ambiguous case with 0/0/0",
            },
            {
                "path": "Detector residual",
                "cases_affected": len(detector_blocked),
                "safe_known_gain": 0,
                "plausible_gain": stage_b_new,
                "risk": f"MEDIUM-HIGH: {stage_b_new} gains span sparse aliases/patterns after V4 closure",
                "complexity": "MEDIUM",
                "generalization": "LOW-MEDIUM: residuals are sparse and mechanism-diverse",
                "gold_dependence": "oracle family assignment uses target nature and is not a production proposal",
                "thresholds": "not evaluated",
                "evidence": f"root-aware immediate upper bound={stage_b_new}; legacy taxonomy reports {legacy_detector_upper}",
            },
            {
                "path": "Parser",
                "cases_affected": stage_c_new,
                "safe_known_gain": 0,
                "plausible_gain": stage_c_new,
                "risk": "MEDIUM: all require joint complete-span + OCR repair and are concentrated in N2",
                "complexity": "LOW-MEDIUM",
                "generalization": "MEDIUM-LOW: three N2 cases, all requiring bounded representation repair",
                "gold_dependence": "family is oracle-only; repaired values come from surface",
                "thresholds": "none; fixed character/separator normalization",
                "evidence": "parser-only current upper bound=0; joint oracle adds three",
            },
            {
                "path": "Legal",
                "cases_affected": len(legal_real),
                "safe_known_gain": legal_e2e_gain if legal_article_safe else 0,
                "plausible_gain": legal_e2e_gain,
                "risk": "HIGH until source-diploma metadata exists; article-only creates false-real",
                "complexity": "HIGH (corpus enrichment), LOW (unsafe article lookup)",
                "generalization": "LOW without verified source identity; article number alone is not legally canonical",
                "gold_dependence": "article-only candidates are gold-free, but its apparent 14/14 real oracle fit is unsafe globally",
                "thresholds": "none",
                "evidence": f"article-only +{article_safety['correct'] - baseline_safety['correct']} but false-real={article_safety['false_real_inventada']} and {article_formal['status']['resolved']} resolved formal FPs",
            },
            {
                "path": "Semantic/contextual retrieval",
                "cases_affected": len(semantic_cases),
                "safe_known_gain": 0,
                "plausible_gain": 0,
                "risk": "HIGH/undefined without a qualifying pool",
                "complexity": "MEDIUM-HIGH",
                "generalization": "UNASSESSED: no genuine semantic case remains",
                "gold_dependence": "not applicable",
                "thresholds": "would be required; no experiment justified",
                "evidence": f"revised semantic pool is empty; all {len(noisy_structural)} reviewed cases are noisy structural references with unique correct structural candidates",
            },
            {
                "path": "Classification final",
                "cases_affected": 225,
                "safe_known_gain": "12 class labels (not IDs)",
                "plausible_gain": "12 selective labels",
                "risk": "HIGH for a complete classifier; LOW for the single observed incomplete-only rule",
                "complexity": "MEDIUM",
                "generalization": "LOW for full status mapping; MEDIUM for the bounded general-reference rule",
                "gold_dependence": "status-majority bounds are gold-fitted; the selective rule is merely observational",
                "thresholds": "none for the bounded rule",
                "evidence": "missing/insufficient mix all classes; one selective family+status rule is 12/12",
            },
        ]

        review_indexes = {
            conflict_item.index,
            *(int(case["gold_index"]) for case in cnj_ambiguous),
            *(item.gold.index for item in current_ambiguous),
            *(int(item["gold_index"]) for item in real_no_match),
        }
        human_review = [audit for audit in unresolved if int(audit["gold_index"]) in review_indexes]
        for audit in human_review:
            reasons: list[str] = []
            if audit["gold_index"] == conflict_item.index:
                reasons.extend(["CNJ/canonical identity conflict", "confirmed gold/corpus inconsistency"])
            if any(int(case["gold_index"]) == audit["gold_index"] for case in cnj_ambiguous):
                reasons.append("multi-ID near-duplicate CNJ")
            if any(item.gold.index == audit["gold_index"] for item in current_ambiguous):
                reasons.append("current resolver ambiguous")
            if any(int(item["gold_index"]) == audit["gold_index"] for item in real_no_match):
                reasons.append("real current no_match")
            audit["human_review_reasons"] = reasons
            if audit["gold_index"] == conflict_item.index:
                audit["human_review_special_evidence"] = {
                    "primary_identity_conflict": conflict
                }

        payload: dict[str, object] = {
            "status": "PASS WITH WARNINGS",
            "baseline": {
                "dataset": {
                    "directory": str(DATASET_DIR.relative_to(PROJECT_ROOT)),
                    "documents": len(texts),
                    "citations": len(gold),
                    "classes": dict(sorted(EXPECTED_CLASSES.items())),
                    "levels": dict(sorted(EXPECTED_LEVELS.items())),
                },
                "database": {
                    "records": sum(database_counts.values()),
                    "nature": database_counts,
                    "sha256": database_hash,
                    "integrity_check": integrity,
                    "case_index": case_index.stats,
                    "document_columns": list(database_columns),
                },
                "detector_v4": detector,
                "end_to_end": baseline_safety,
                "formal_false_positives": baseline_formal_fp,
                "tests_expected": "65/65 (executed outside this script)",
            },
            "unresolved_real": unresolved,
            "first_blockers": first_blockers,
            "root_causes": roots,
            "root_cause_matrix": root_matrix,
            "actionable_components": actionable,
            "recovery_ladder": recovery_ladder,
            "candidate_set_entropy": {bucket: entropy[bucket] for bucket in ("0", "1", "2-5", ">5")},
            "coverage_dimensions": {
                "detection": {
                    "matched_gold": detector["TP"],
                    "gold_total": len(gold),
                    "formal_predictions": detector["predictions"],
                    "by_gold_class": classification["detected_by_gold"],
                },
                "structural_parse": {
                    "identity_preserved_matched": detector["identity_preservation"],
                    "matched_gold": detector["TP"],
                    "note": "parser equivalence is family-specific field equivalence against full-span parsing, not canonical resolution",
                },
                "canonical_id": {
                    "correct_real": baseline_safety["correct"],
                    "real_total": len(real_evaluations),
                    "safety": {
                        key: baseline_safety[key]
                        for key in ("wrong_unique_real", "false_real_inventada", "false_real_incompleta")
                    },
                },
                "classifiability": classification["simple_detection_gated_mapping"],
            },
            "detector_residual": detector_residual,
            "parser": parser_audit,
            "cnj": {
                "strict_gold_universe": {
                    "total": len(strict_cnj_items),
                    "real": sum(item.classification == "real" for item in strict_cnj_items),
                    "inventada": sum(item.classification == "inventada" for item in strict_cnj_items),
                    "incompleta": sum(item.classification == "incompleta" for item in strict_cnj_items),
                    "physical_number_overlap_real": sum(
                        item.classification == "real"
                        and any(
                            prediction.document_id == item.document_id
                            and iou(item.start, item.end, prediction.start, prediction.end) > 0
                            and prediction.family == "processo_cnj"
                            for prediction in predictions
                        )
                        for item in strict_cnj_items
                    ),
                    "formal_matched_real": sum(
                        item.classification == "real" and by_gold_eval[item.index].prediction is not None
                        for item in strict_cnj_items
                    ),
                    "formal_matched_total": sum(
                        by_gold_eval[item.index].prediction is not None for item in strict_cnj_items
                    ),
                    "physical_process_cnj_overlap_total": sum(
                        any(
                            prediction.document_id == item.document_id
                            and iou(item.start, item.end, prediction.start, prediction.end) > 0
                            and prediction.family == "processo_cnj"
                            for prediction in predictions
                        )
                        for item in strict_cnj_items
                    ),
                    "current_parser_full_span_normalized": sum(
                        bool(
                            parser.parse(
                                CitationCandidate(
                                    0,
                                    len(texts[item.document_id][item.start : item.end]),
                                    texts[item.document_id][item.start : item.end],
                                    "cnj_full_span_control",
                                    "processo_cnj",
                                )
                            ).data.get("numero_normalizado")
                        )
                        for item in strict_cnj_items
                    ),
                    "exact_candidate_sets": sum(
                        bool(cnj_queries[item.index].candidate_ids) for item in strict_cnj_items
                    ),
                },
                "broad_relevant_universe": {
                    "total": len(broad_cnj_items),
                    "real": sum(item.classification == "real" for item in broad_cnj_items),
                    "inventada": sum(item.classification == "inventada" for item in broad_cnj_items),
                    "additional_degraded_surfaces": [
                        item.index for item in broad_cnj_items if item not in strict_cnj_items
                    ],
                    "remaining_real": sum(
                        item.classification == "real" and by_gold_eval[item.index].blocker != "resolved_correct"
                        for item in broad_cnj_items
                    ),
                    "formal_matched_total": sum(
                        by_gold_eval[item.index].prediction is not None for item in broad_cnj_items
                    ),
                    "oracle_structural_number_available": sum(
                        bool(broad_cnj_queries[item.index].number) for item in broad_cnj_items
                    ),
                    "exact_candidate_sets": sum(
                        bool(broad_cnj_queries[item.index].candidate_ids) for item in broad_cnj_items
                    ),
                },
                "cases": cnj_cases,
                "broad_cases": broad_cnj_cases,
                "strategy_matrix_oracle_strict": cnj_matrix,
                "strategy_matrix_oracle_broad": broad_cnj_matrix,
                "strategy_matrix_current_detector": current_cnj["matrix"],
                "resolved_formal_fp_details": current_cnj["resolved_formal_fp_details"],
                "conflict": conflict,
                "ambiguous": cnj_ambiguous,
                "gold_safe_vs_corpus_consistent": {
                    "cnj_exact": "corpus-consistent but not gold-safe",
                    "non_tst_primary_exact": "gold-safe and corpus-consistent; conservative +12",
                    "primary_class_guard": "gold-safe and corpus-consistent by abstention; +14",
                    "body_fallback": "closed-corpus gold-safe but corpus-inconsistent and not promotion-ready",
                },
                "verdict": "SAFE_PATH_EXISTS",
                "verdict_reason": "a runtime-available TST procedural-class guard removes the sole wrong unique without sacrificing the +14 current gain",
            },
            "legal": {
                "remaining_real": len(legal_real),
                "cases": legal_cases,
                "corpus_records": len(article_index),
                "source_diploma_metadata_records": None,
                "structured_source_identity_fields_available": list(structured_legal_identity_fields),
                "source_identity_availability": (
                    "AVAILABLE" if structured_legal_identity_fields else "NOT_AVAILABLE_IN_SCHEMA"
                ),
                "article_keys_unique": all(len(ids) == 1 for ids in article_index.values()),
                "strategy_matrix": [
                    {
                        "strategy": "article_only",
                        "evaluation_status": "EVALUATED_UNSAFE",
                        "support": len(legal_cases),
                        "oracle_correct": legal_oracle_correct,
                        "e2e_gain": legal_e2e_gain,
                        "wrong_unique_real": article_safety["wrong_unique_real"],
                        "false_real": article_safety["false_real_inventada"] + article_safety["false_real_incompleta"],
                        "formal_fp_resolved": article_formal["status"]["resolved"],
                        "formal_fp_near_invented": article_formal_near_invented,
                    },
                    {
                        "strategy": "article_diploma",
                        "evaluation_status": "NOT_EVALUABLE",
                        "support": 0,
                        "oracle_correct": None,
                        "e2e_gain": None,
                        "wrong_unique_real": None,
                        "false_real": None,
                        "formal_fp_resolved": None,
                        "reason": "source diploma is absent from corpus metadata",
                    },
                    {
                        "strategy": "article_law_number_year",
                        "evaluation_status": "NOT_EVALUABLE",
                        "support": 0,
                        "oracle_correct": None,
                        "e2e_gain": None,
                        "wrong_unique_real": None,
                        "false_real": None,
                        "formal_fp_resolved": None,
                        "reason": "body law mentions are amendments/references, not primary identity",
                    },
                ],
                "unsafe_examples": [
                    {
                        "gold_index": item.gold.index,
                        "documento": item.gold.document_id,
                        "trecho": item.gold.text,
                        "resolved_id": item.current_result.id_canonico,
                    }
                    for item in unsafe_legal
                ],
                "article_only_resolved_formal_fp_details": article_formal_details,
                "oracle_potential": legal_oracle_correct,
                "current_detector_plausible_gain_with_verified_identity": legal_e2e_gain,
                "safe_known_gain": legal_e2e_gain if legal_article_safe else 0,
                "roi": "LOW until canonical legal source metadata is added",
            },
            "semantic": {
                "total_real_cases": len(semantic_cases),
                "families": dict(Counter(str(item["oracle_family"]) for item in semantic_cases)),
                "retrieval_readiness": {level: semantic_readiness[level] for level in ("HIGH", "MEDIUM", "LOW")},
                "cases": semantic_cases,
                "reclassified_noisy_structural_pool": {
                    "total": len(noisy_structural),
                    "N1": sum(item["nivel"] == "N1" for item in noisy_structural),
                    "N2": sum(item["nivel"] == "N2" for item in noisy_structural),
                    "documents": len({str(item["documento"]) for item in noisy_structural}),
                    "gold_indexes": [item["gold_index"] for item in noisy_structural],
                    "unique_candidate_correct": noisy_unique_correct,
                    "structurally_resolvable": sum(
                        bool(item["recovery_ladder"]["structurally_resolvable"])
                        for item in noisy_structural
                    ),
                    "evidence": "surface-derived normalized structural queries, evaluated against gold only after candidate construction",
                },
                "retrieval_experiment": {
                    "executed": False,
                    "reason": "no genuinely semantic/contextual unresolved real case remains after root-cause correction",
                    "recommendation": "do not run BM25/FTS until a non-structural pool is demonstrated",
                },
                "future_backoff_order_if_new_pool_appears": [
                    "deterministic metadata filtering",
                    "local SQLite FTS/BM25 candidate-recall experiment",
                    "TF-IDF only if FTS diagnostics are inadequate",
                    "embeddings only with demonstrated lexical failure and sufficient evaluation support",
                ],
                "genuine_semantic_metadata_filterable_before_retrieval": sum(
                    bool(item["case_index"]["class_candidate_ids"])
                    or bool(item["case_index"]["tribunal_candidate_ids"])
                    for item in semantic_cases
                ),
                "overall_structural_cases_misclassified_as_semantic_before_review": len(noisy_structural),
            },
            "ambiguous": {
                "current": len(current_ambiguous),
                "oracle_or_experimental_total": len(all_ambiguous),
                "class_filter_safe_gain": ambiguity_safety["correct"] - baseline_safety["correct"],
                "class_filter_safety": {
                    key: ambiguity_safety[key]
                    for key in ("wrong_unique_real", "false_real_inventada", "false_real_incompleta")
                },
                "cases": all_ambiguous,
            },
            "real_no_match": real_no_match,
            "gold_issues": {
                "confirmed": [conflict],
                "probable": [],
                "possible": [],
                "counts": {"confirmed": 1, "probable": 0, "possible": 0},
                "note": "two duplicate candidate sets are corpus ambiguity, not labeled as gold inconsistency",
            },
            "corpus_quality": {
                "primary_vs_body_mention_conflicts": 1,
                "near_duplicate_multi_id_cases": 2,
                "missing_sumula_number_metadata": 2,
                "missing_legal_source_identity_cases": 14,
                "type_mismatch": 0,
                "missing_canonical_record": 0,
            },
            "classification_readiness": classification,
            "invented_gold_audit": invented_audit,
            "incomplete_gold_control": incomplete_control,
            "roadmap": roadmap,
            "diminishing_returns": {
                "Detector": "nearing saturation",
                "Parser": "saturated for current spans",
                "Resolver structural": "productive",
            },
            "human_review": {
                "artifact": str(HTML_PATH.relative_to(PROJECT_ROOT)),
                "case_count": len(human_review),
                "categories": {
                    "CNJ/canonical conflicts": 3,
                    "real no_match": len(real_no_match),
                    "ambiguous multi-ID": len(all_ambiguous),
                    "semantic HIGH": 0,
                    "possible gold inconsistencies": 1,
                },
                "cases": human_review,
            },
            "recommendation": {
                "primary": "A) CNJ / canonical identity",
                "now": "Implement in a separate task a CNJ primary-identity resolver with TST procedural-class guard, explicit abstention, and overlap/duplicate arbitration.",
                "then": "Add the general primary-class filter for the one STJ ambiguity, then design a bounded noisy-structural front-end experiment for the 15 OCR/alias cases.",
                "later_if_needed": "Enrich sumula/legal corpus identity and add only the observed selective incomplete rule; reconsider retrieval only if a genuine non-structural real pool appears.",
                "why": [
                    "+14 canonical IDs is the largest promotion-grade safe observed gain.",
                    "The class guard preserves 0 wrong-unique and 0 false-real on all gold classes.",
                    "The only cnj_exact wrong is avoidable by deterministic abstention.",
                    "Twenty-seven unresolved real cases have a plausible gold-consistent unique CNJ path under an oracle front end.",
                    "The two remaining CNJ multi-ID sets are near-duplicates and can remain ambiguous safely.",
                    "Tribunal and nature are redundant; class is the decisive runtime signal.",
                    "Legal article-only is unsafe and the corpus lacks source-diploma identity.",
                    "Parser-only immediate gain is zero and Detector V4 is near saturation.",
                    "The revised semantic pool is zero, so BM25/FTS would be premature.",
                    "Full three-way classification still confuses unsupported real citations with incomplete ones.",
                ],
            },
            "blockers": [
                "two TST duplicate pairs have no runtime-distinguishing metadata and must remain ambiguous",
                "fourteen legal occurrences and two sumulas lack verified canonical identity metadata",
                "one confirmed gold/source-document versus primary-identity conflict cannot be structurally resolved to the gold ID",
            ],
            "warnings": [
                "the safe CNJ experiment resolves six formally false-positive nested spans; production promotion needs overlap arbitration/deduplication",
                "the body-mention fallback passes closed-corpus safety but changes identity semantics and is not promotion-ready",
                "article-only legal lookup creates one matched false-real and seventeen resolved formal false positives",
                "the gold-fitted 67.6% status oracle is not runtime-achievable; the simple detection-gated mapping is only 84/225 and emits 84 formal false positives",
            ],
            "determinism": {"required_runs": 3, "verified_in_script": False, "stable_projection": "entire JSON payload"},
        }

    performance = {
        "pipeline_seconds": pipeline_seconds,
        "audit_seconds": time.perf_counter() - started,
        "seconds_per_document": pipeline_seconds / len(texts),
    }
    return payload, performance


def html_table(headers: Sequence[str], rows: Iterable[Sequence[object]]) -> str:
    head = "".join(f"<th>{html.escape(str(header))}</th>" for header in headers)
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in row) + "</tr>"
        for row in rows
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def write_html(payload: Mapping[str, object]) -> None:
    review = payload["human_review"]
    cases = review["cases"]
    cards: list[str] = []
    for case in cases:
        target = case["gold_only_evidence"]["target_record"]
        candidates = case["case_index"]["identity_candidate_ids"]
        candidate_records = case["case_index"]["identity_candidate_records"]
        related = case["detector_candidates_relevantes"]
        special = case.get("human_review_special_evidence")
        cards.append(
            "<article>"
            f"<h2>gold {case['gold_index']} · {html.escape(str(case['documento']))}</h2>"
            f"<p><strong>Review:</strong> {html.escape(', '.join(case.get('human_review_reasons', [])))}</p>"
            f"<p><strong>First blocker:</strong> {html.escape(str(case['first_blocker']))} · "
            f"<strong>root:</strong> {html.escape(str(case['root_cause']))} · "
            f"<strong>action:</strong> {html.escape(str(case['actionable_component']))}</p>"
            f"<pre>{html.escape(str(case['gold_trecho']))}</pre>"
            f"<details open><summary>Context ±200</summary><pre>{html.escape(str(case['runtime_available_evidence']['context_200']))}</pre></details>"
            f"<details><summary>Current ParsedCitation / ResolutionResult</summary><pre>{html.escape(json.dumps({'parsed': case['current_parsed'], 'result': case['current_resolution']}, ensure_ascii=False, indent=2))}</pre></details>"
            f"<details><summary>Oracle / candidate IDs</summary><pre>{html.escape(json.dumps({'oracle': case['oracle_parsed'], 'resolution': case['oracle_resolution'], 'candidate_ids': candidates}, ensure_ascii=False, indent=2))}</pre></details>"
            f"<details><summary>Candidate corpus metadata/snippets</summary><pre>{html.escape(json.dumps(candidate_records, ensure_ascii=False, indent=2))}</pre></details>"
            f"<details><summary>Target corpus metadata</summary><pre>{html.escape(json.dumps(target, ensure_ascii=False, indent=2))}</pre></details>"
            f"<details><summary>Special review evidence</summary><pre>{html.escape(json.dumps(special, ensure_ascii=False, indent=2))}</pre></details>"
            f"<details><summary>Relevant detector candidates ({len(related)})</summary><pre>{html.escape(json.dumps(related, ensure_ascii=False, indent=2))}</pre></details>"
            "</article>"
        )
    category_rows = [(key, value) for key, value in review["categories"].items()]
    document = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Post-Detector V4 Deep Audit — Human Review</title>
<style>
body{{font:15px/1.45 system-ui,sans-serif;margin:2rem auto;max-width:1180px;padding:0 1rem;color:#18202a}}
h1{{margin-bottom:.25rem}} article{{border:1px solid #ccd4df;border-radius:10px;padding:1rem;margin:1.25rem 0}}
pre{{white-space:pre-wrap;overflow-wrap:anywhere;background:#f5f7fa;padding:.75rem;border-radius:6px}}
table{{border-collapse:collapse;margin:1rem 0}} th,td{{border:1px solid #ccd4df;padding:.45rem .65rem;text-align:left}}
th{{background:#edf2f7}} summary{{cursor:pointer;font-weight:600}} .warning{{background:#fff4d6;padding:.75rem;border-left:4px solid #d99b00}}
</style></head><body>
<h1>Post-Detector V4 Deep Audit</h1>
<p>Human-review set: {review['case_count']} casos (união deduplicada das categorias obrigatórias).</p>
<p class="warning">Este artefato usa gold para auditoria. Nenhuma decisão experimental aqui integra runtime de produção.</p>
{html_table(['categoria', 'casos'], category_rows)}
{''.join(cards)}
</body></html>"""
    HTML_PATH.parent.mkdir(parents=True, exist_ok=True)
    HTML_PATH.write_text(document, encoding="utf-8")


def print_table(headers: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    string_rows = [[str(value) for value in row] for row in rows]
    widths = [len(str(header)) for header in headers]
    for row in string_rows:
        for index, value in enumerate(row):
            widths[index] = max(widths[index], len(value))
    print(" | ".join(str(header).ljust(widths[index]) for index, header in enumerate(headers)))
    print("-+-".join("-" * width for width in widths))
    for row in string_rows:
        print(" | ".join(value.ljust(widths[index]) for index, value in enumerate(row)))


def print_summary(payload: Mapping[str, object], performance: Mapping[str, float]) -> None:
    baseline = payload["baseline"]
    detector = baseline["detector_v4"]
    safety = baseline["end_to_end"]
    print("=== BASELINE ===")
    print(
        f"dataset=26/225; DB=1016 (998/13/5); Detector V4="
        f"{detector['predictions']}/{detector['TP']}/{detector['FP']}/{detector['FN']}, exact={detector['exact']}"
    )
    print(
        f"E2E={safety['correct']}/96; safety="
        f"{safety['wrong_unique_real']}/{safety['false_real_inventada']}/{safety['false_real_incompleta']}; "
        f"formal FPs insufficient={baseline['formal_false_positives']['status']['insufficient']}"
    )

    print("\n=== 55 UNRESOLVED REAL ===")
    print(f"derived={len(payload['unresolved_real'])}; JSON={JSON_PATH.relative_to(PROJECT_ROOT)}")

    print("\n=== FIRST BLOCKERS ===")
    print_table(["first blocker", "count", "%"], ((row["name"], row["cases"], row["percent"]) for row in payload["first_blockers"]))

    print("\n=== ROOT CAUSES ===")
    print_table(["root cause", "count", "%"], ((row["name"], row["cases"], row["percent"]) for row in payload["root_causes"]))
    print("actionable components:")
    print_table(["component", "count", "%"], ((row["name"], row["cases"], row["percent"]) for row in payload["actionable_components"]))

    print("\n=== RECOVERY LADDER ===")
    print_table(["stage", "newly recovered", "cumulative"], ((row["stage"], row["newly_recovered"], row["cumulative"]) for row in payload["recovery_ladder"]))
    print(
        f"detector immediate upper={payload['detector_residual']['detector_only_immediate_upper_bound_post_v4']}; "
        f"parser-only immediate={payload['parser']['parser_only_immediate_upper_bound']}"
    )

    print("\n=== CNJ DEEP AUDIT ===")
    cnj = payload["cnj"]
    print(
        f"strict={cnj['strict_gold_universe']['total']} "
        f"(real={cnj['strict_gold_universe']['real']}, invented={cnj['strict_gold_universe']['inventada']}); "
        f"broad relevant={cnj['broad_relevant_universe']['total']}; remaining real={cnj['broad_relevant_universe']['remaining_real']}"
    )
    print(
        "strict coverage: "
        f"formal matched={cnj['strict_gold_universe']['formal_matched_total']}, "
        f"physical CNJ overlap={cnj['strict_gold_universe']['physical_process_cnj_overlap_total']}, "
        f"full-span parsed={cnj['strict_gold_universe']['current_parser_full_span_normalized']}, "
        f"exact candidate sets={cnj['strict_gold_universe']['exact_candidate_sets']}"
    )
    print_table(
        ["strategy", "support", "correct", "wrong", "ambiguous", "no_match", "false-real"],
        (
            (
                row["strategy"], row["support"], row["correct_unique"], row["wrong_unique"],
                row["ambiguous"], row["no_match"], row["false_real"],
            )
            for row in cnj["strategy_matrix_oracle_strict"]
        ),
    )
    print("broad oracle strategies (includes eight degraded surfaces):")
    print_table(
        ["strategy", "support", "correct", "wrong", "ambiguous", "no_match", "false-real"],
        (
            (
                row["strategy"], row["support"], row["correct_unique"], row["wrong_unique"],
                row["ambiguous"], row["no_match"], row["false_real"],
            )
            for row in cnj["strategy_matrix_oracle_broad"]
        ),
    )
    print("current-detector E2E strategies:")
    print_table(
        ["strategy", "gain", "wrong", "false inv", "false inc", "resolved formal FP"],
        (
            (
                row["strategy"], row["gain"], row["wrong_unique_real"], row["false_real_inventada"],
                row["false_real_incompleta"], row["formal_fp_resolved"],
            )
            for row in cnj["strategy_matrix_current_detector"]
        ),
    )
    print(f"conflict gold={cnj['conflict']['gold_id']} primary={cnj['conflict']['primary_exact_id']}; verdict={cnj['verdict']}")

    print("\n=== LEGAL AUDIT ===")
    legal = payload["legal"]
    print(
        f"remaining={legal['remaining_real']}; oracle={legal['oracle_potential']}; "
        f"current plausible={legal['current_detector_plausible_gain_with_verified_identity']}; safe known={legal['safe_known_gain']}"
    )
    print_table(
        ["strategy", "oracle", "E2E gain", "false-real", "formal FP resolved"],
        (
            (row["strategy"], row["oracle_correct"], row["e2e_gain"], row["false_real"], row["formal_fp_resolved"])
            for row in legal["strategy_matrix"]
        ),
    )

    print("\n=== SEMANTIC/CONTEXTUAL AUDIT ===")
    semantic = payload["semantic"]
    print(
        f"genuine total={semantic['total_real_cases']}; readiness={semantic['retrieval_readiness']}; "
        f"reclassified noisy structural={semantic['reclassified_noisy_structural_pool']['total']}"
    )
    print(f"retrieval experiment: NOT RUN — {semantic['retrieval_experiment']['reason']}")

    print("\n=== AMBIGUOUS / NO-MATCH ===")
    print(
        f"ambiguous current={payload['ambiguous']['current']}; all structural={payload['ambiguous']['oracle_or_experimental_total']}; "
        f"class-filter safe gain={payload['ambiguous']['class_filter_safe_gain']}; real no-match={len(payload['real_no_match'])}"
    )
    for item in payload["real_no_match"]:
        print(f"  gold {item['gold_index']} {item['documento']}: {item['reason']}")

    print("\n=== GOLD / CORPUS ISSUES ===")
    print(f"confirmed/probable/possible={payload['gold_issues']['counts']}")

    print("\n=== CLASSIFICATION READINESS ===")
    matrix = payload["classification_readiness"]["status_by_gold"]
    print_table(
        ["gold", *STATUS_ORDER],
        ((classification, *(matrix[classification][status] for status in STATUS_ORDER)) for classification in ("real", "inventada", "incompleta")),
    )
    opportunity = payload["classification_readiness"]["classification_only_opportunity"]
    mapping = payload["classification_readiness"]["simple_detection_gated_mapping"]
    invented = payload["invented_gold_audit"]
    print(
        f"readiness={payload['classification_readiness']['readiness']}; selective opportunity="
        f"{opportunity['observed_correct']}/{opportunity['support']} labels; "
        f"detection-gated={mapping['correct']}/{mapping['gold_total']} gold "
        f"({mapping['correct']}/{mapping['detector_predictions_emitted']} emitted predictions)"
    )
    print(
        f"invented audit: current no_match/insufficient/missing="
        f"{invented['current_status']['no_match']}/{invented['current_status']['insufficient']}/"
        f"{invented['current_status']['missing']}; full-span surface oracle="
        f"{invented['surface_oracle_current_parser_resolver']['status']['no_match']}/"
        f"{invented['surface_oracle_current_parser_resolver']['status']['insufficient']}"
    )

    print("\n=== ROADMAP ===")
    print_table(
        ["path", "affected", "safe known", "plausible", "risk", "complexity"],
        ((row["path"], row["cases_affected"], row["safe_known_gain"], row["plausible_gain"], row["risk"], row["complexity"]) for row in payload["roadmap"]),
    )

    print("\n=== RECOMMENDATION ===")
    recommendation = payload["recommendation"]
    print(recommendation["primary"])
    print(f"NOW: {recommendation['now']}")
    print(f"THEN: {recommendation['then']}")
    print(f"LATER IF NEEDED: {recommendation['later_if_needed']}")
    print(
        f"human review={HTML_PATH.relative_to(PROJECT_ROOT)} ({payload['human_review']['case_count']} cases); "
        f"determinism={payload['determinism']['runs_verified']}/{payload['determinism']['required_runs']}; "
        f"performance pipeline={performance['pipeline_seconds']:.3f}s audit={performance['audit_seconds']:.3f}s"
    )


def main(argv: Sequence[str] | None = None) -> int:
    single_run = "--single-run" in (argv if argv is not None else sys.argv[1:])
    runs = 1 if single_run else 3
    try:
        payloads: list[dict[str, object]] = []
        performances: list[dict[str, float]] = []
        for _ in range(runs):
            payload, performance = build_audit()
            payloads.append(payload)
            performances.append(performance)
        reference = json.dumps(payloads[0], ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        for payload in payloads[1:]:
            current = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            if current != reference:
                raise BaselineError("structured output is not deterministic across runs")
        payload = payloads[0]
        payload["determinism"] = {
            "required_runs": 3,
            "runs_verified": runs,
            "verified": runs == 3,
            "stable_projection": "entire JSON payload excluding runtime timings",
        }
        JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
        JSON_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        write_html(payload)
        print_summary(payload, performances[0])
        return 0
    except BaselineError as exc:
        print(f"BASELINE ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
